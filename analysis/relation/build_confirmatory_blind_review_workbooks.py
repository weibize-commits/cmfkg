from __future__ import annotations

import hashlib
import json
import random
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[2]
EXPERIMENT_DIR = ROOT / "paper_revision" / "relation_improvement"
DATA_DIR = EXPERIMENT_DIR / "data"
REVIEW_DIR = EXPERIMENT_DIR / "human_revalidation"
SEED = 20260731


RELATION_DEFINITIONS = {
    "hasNature": "食材具有某种食性（寒、凉、平、温或热）",
    "hasFlavor": "食材具有某种食味（酸、苦、甘、辛或咸）",
    "entersmeridian": "食材归入某经",
    "hasEffect": "食材具有某种功效",
    "treats": "食材治疗或缓解某病症",
    "containsIngredient": "食疗方含有某食材",
    "usesCookingMethod": "食疗方使用某烹饪方法",
    "treatedByRecipe": "病症由某食疗方治疗或缓解",
    "contraindicates": "食材存在某种食用禁忌",
    "affectsOrgan": "病症涉及或影响某脏腑",
    "correspondsToOrgan": "归经对应某脏腑",
    "recordedInText": "实体在某古籍中有记载",
    "authoredBy": "古籍由某医家所著",
    "recommendsIngredient": "养生原则推荐某食材",
}


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def review_id(candidate: dict[str, Any]) -> str:
    head = candidate["head"]
    tail = candidate["tail"]
    value = (
        str(candidate["sample_id"]),
        int(head["start_char"]),
        int(head["end_char"]),
        "".join(str(head["text"]).split()),
        str(head["type"]),
        str(candidate["relation"]),
        int(tail["start_char"]),
        int(tail["end_char"]),
        "".join(str(tail["text"]).split()),
        str(tail["type"]),
    )
    payload = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return "REL-" + hashlib.sha256(payload.encode("utf-8")).hexdigest().upper()[:16]


def candidate_rows() -> list[list[Any]]:
    candidates = read_jsonl(DATA_DIR / "test_candidates.jsonl")
    passages = read_jsonl(DATA_DIR / "test_passages.jsonl")
    passage_by_id = {
        str(row["sample_id"]): re.sub(r"\s+", " ", str(row["frozen_text"]))
        for row in passages
    }
    rows: list[list[Any]] = []
    for candidate in candidates:
        head = candidate["head"]
        tail = candidate["tail"]
        relation = str(candidate["relation"])
        rows.append(
            [
                review_id(candidate),
                str(candidate["sample_id"]),
                passage_by_id[str(candidate["sample_id"])],
                re.sub(r"\s+", " ", str(candidate.get("evidence_text") or "")),
                str(head["text"]),
                str(head["type"]),
                int(head["start_char"]),
                int(head["end_char"]),
                relation,
                RELATION_DEFINITIONS.get(relation, relation),
                str(tail["text"]),
                str(tail["type"]),
                int(tail["start_char"]),
                int(tail["end_char"]),
            ]
        )
    return rows


def style_form_sheet(sheet, decision_column: str, max_row: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = {
        "A": 22,
        "B": 14,
        "C": 100,
        "D": 70,
        "E": 24,
        "F": 16,
        "G": 12,
        "H": 12,
        "I": 24,
        "J": 42,
        "K": 24,
        "L": 16,
        "M": 12,
        "N": 12,
        "O": 24,
        "P": 50,
        "Q": 45,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    validation = DataValidation(
        type="list",
        formula1='"present,absent,insufficient_context,entity_or_type_error"',
        allow_blank=True,
    )
    sheet.add_data_validation(validation)
    validation.add(f"{decision_column}2:{decision_column}{max(2, max_row)}")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def add_instructions(workbook: Workbook, role: str) -> None:
    sheet = workbook.active
    sheet.title = "00_Instructions"
    rows = [
        ("角色", role),
        ("任务单位", "每一行是一个固定方向、固定实体边界和固定关系类型的候选。"),
        (
            "present",
            "当前段落直接表达该关系，或条目标题、属性列举、配方原料/制法结构明确支持该关系。",
        ),
        (
            "absent",
            "只是共现、关系方向错误、尾实体属于其他条目，或当前段落没有足够证据支持。",
        ),
        (
            "insufficient_context",
            "段落截断或OCR损坏导致无法判断，但实体边界和类型本身未必错误。",
        ),
        (
            "entity_or_type_error",
            "头尾实体边界或实体类型明显错误，以致给定关系候选无法有效判断。",
        ),
        (
            "证据要求",
            "判present时必须填写不超过50字的原文证据；其他判断应在notes中简要说明原因。",
        ),
        (
            "盲法要求",
            "不得查看另一位专家、现有金标准、模型预测或模型置信度。不得使用段落外知识补充关系。",
        ),
    ]
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 110
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_expert_workbook(
    rows: list[list[Any]],
    expert_name: str,
    seed: int,
    output_path: Path,
) -> None:
    workbook = Workbook()
    add_instructions(workbook, expert_name)
    form = workbook.create_sheet("01_Blind_Ratings")
    headers = [
        "review_id",
        "sample_id",
        "full_passage_text",
        "local_evidence_window",
        "head_text",
        "head_type",
        "head_start",
        "head_end",
        "relation",
        "relation_definition",
        "tail_text",
        "tail_type",
        "tail_start",
        "tail_end",
        "relation_decision",
        "evidence_quote",
        "notes",
    ]
    form.append(headers)
    shuffled = rows[:]
    random.Random(seed).shuffle(shuffled)
    for row in shuffled:
        form.append(row + ["", "", ""])
    style_form_sheet(form, "O", form.max_row)
    workbook.save(output_path)


def build_coordinator_workbook(rows: list[list[Any]], output_path: Path) -> None:
    workbook = Workbook()
    add_instructions(workbook, "协调员与第三位仲裁专家")
    form = workbook.create_sheet("01_Merge_and_Adjudicate")
    headers = [
        "review_id",
        "sample_id",
        "full_passage_text",
        "local_evidence_window",
        "head_text",
        "head_type",
        "head_start",
        "head_end",
        "relation",
        "relation_definition",
        "tail_text",
        "tail_type",
        "tail_start",
        "tail_end",
        "expert1_decision",
        "expert1_evidence",
        "expert1_notes",
        "expert2_decision",
        "expert2_evidence",
        "expert2_notes",
        "agreement_status",
        "adjudicated_decision",
        "adjudication_evidence_and_reason",
    ]
    form.append(headers)
    for row in sorted(rows, key=lambda value: value[0]):
        form.append(row + ["", "", "", "", "", "", "", "", ""])
    style_form_sheet(form, "V", form.max_row)
    for column, width in {"R": 24, "S": 50, "T": 45, "U": 20, "V": 24, "W": 60}.items():
        form.column_dimensions[column].width = width
    workbook.save(output_path)


def main() -> None:
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    rows = candidate_rows()
    build_expert_workbook(
        rows,
        "独立专家1",
        SEED + 1,
        REVIEW_DIR / "01_RelationConfirmatory_Expert1_Blind.xlsx",
    )
    build_expert_workbook(
        rows,
        "独立专家2",
        SEED + 2,
        REVIEW_DIR / "02_RelationConfirmatory_Expert2_Blind.xlsx",
    )
    build_coordinator_workbook(
        rows,
        REVIEW_DIR / "03_RelationConfirmatory_Coordinator_Adjudication.xlsx",
    )
    print(f"candidate_rows={len(rows)}")
    for path in sorted(REVIEW_DIR.glob("*.xlsx")):
        print(path)


if __name__ == "__main__":
    main()
