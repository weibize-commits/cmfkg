from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
EXPERIMENT_DIR = ROOT / "paper_revision" / "relation_improvement"
DATA_DIR = EXPERIMENT_DIR / "data"
OUTPUT_DIR = EXPERIMENT_DIR / "outputs"


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def key(record: dict[str, Any]) -> tuple[Any, ...]:
    clean = lambda value: "".join(str(value or "").split())
    head = record["head"]
    tail = record["tail"]
    return (
        str(record["sample_id"]),
        int(head["start_char"]),
        int(head["end_char"]),
        clean(head["text"]),
        str(head["type"]),
        str(record["relation"]),
        int(tail["start_char"]),
        int(tail["end_char"]),
        clean(tail["text"]),
        str(tail["type"]),
    )


def digest(record: dict[str, Any]) -> str:
    clean = lambda value: re.sub(r"\s+", " ", str(value or "")).strip()
    head = record["head"]
    tail = record["tail"]
    value = (
        str(record["sample_id"]),
        int(head["start_char"]),
        int(head["end_char"]),
        clean(head["text"]),
        str(head["type"]),
        str(record["relation"]),
        int(tail["start_char"]),
        int(tail["end_char"]),
        clean(tail["text"]),
        str(tail["type"]),
    )
    payload = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest().upper()[:20]


def main() -> None:
    gold = read_jsonl(DATA_DIR / "dev_gold.jsonl")
    passages = read_jsonl(DATA_DIR / "dev_passages.jsonl")
    passage_by_id = {
        str(record["sample_id"]): re.sub(
            r"\s+",
            " ",
            str(record.get("frozen_text") or ""),
        )
        for record in passages
    }
    decisions = read_jsonl(OUTPUT_DIR / "deepseek_triage_dev_decisions.jsonl")
    decision_by_digest = {
        str(record["candidate_digest"]): record for record in decisions
    }
    rows: list[list[Any]] = []
    for record in gold:
        decision = decision_by_digest.get(digest(record))
        if decision is None or decision.get("label") != "unsupported":
            continue
        confidence = float(decision.get("confidence") or 0.0)
        if confidence < 0.85:
            continue
        head = record["head"]
        tail = record["tail"]
        priority = "P1"
        if (
            len(str(head["text"])) > 12
            or len(str(tail["text"])) > 12
            or str(head["text"]) == str(tail["text"])
        ):
            priority = "P0"
        rows.append(
            [
                priority,
                str(record["sample_id"]),
                passage_by_id.get(str(record["sample_id"]), ""),
                str(head["text"]),
                str(head["type"]),
                int(head["start_char"]),
                int(head["end_char"]),
                str(record["relation"]),
                str(tail["text"]),
                str(tail["type"]),
                int(tail["start_char"]),
                int(tail["end_char"]),
                re.sub(r"\s+", " ", str(record.get("evidence_text") or ""))[:800],
                str(record.get("source") or ""),
                str(decision.get("label") or ""),
                confidence,
                str(decision.get("evidence") or ""),
                "",
                "",
                "",
            ]
        )
    rows.sort(key=lambda row: (row[0], row[1], row[7], row[3], row[8]))

    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "00_Instructions"
    instructions.append(["用途", "开发集金标准优先复核，不把模型判断当作真值。"])
    instructions.append(
        [
            "入选标准",
            "现有专家金标准为阳性，但DeepSeek在新提示词下以≥0.85置信度判为unsupported。",
        ]
    )
    instructions.append(
        [
            "复核要求",
            "两位专家独立查看原文、实体边界、实体类型、关系方向和证据；分歧由第三位专家仲裁。",
        ]
    )
    instructions.append(
        [
            "review_decision",
            "填写 keep_gold / remove_relation / correct_entity / correct_relation / insufficient_context。",
        ]
    )
    instructions.append(
        [
            "风险说明",
            "模型仅用于筛选可疑项。任何金标准修改必须以人工复核记录为依据。",
        ]
    )
    instructions.column_dimensions["A"].width = 22
    instructions.column_dimensions["B"].width = 110
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    review = workbook.create_sheet("01_Priority_Review")
    headers = [
        "priority",
        "sample_id",
        "full_passage_text",
        "head_text",
        "head_type",
        "head_start",
        "head_end",
        "relation",
        "tail_text",
        "tail_type",
        "tail_start",
        "tail_end",
        "gold_evidence",
        "gold_source",
        "screening_label",
        "screening_confidence",
        "screening_reason",
        "expert1_decision",
        "expert2_decision",
        "adjudicated_decision_and_note",
    ]
    review.append(headers)
    for row in rows:
        review.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in review[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = {
        "A": 10,
        "B": 14,
        "C": 100,
        "D": 24,
        "E": 16,
        "H": 24,
        "I": 24,
        "J": 16,
        "M": 80,
        "N": 24,
        "O": 18,
        "P": 18,
        "Q": 60,
        "R": 22,
        "S": 22,
        "T": 45,
    }
    for column, width in widths.items():
        review.column_dimensions[column].width = width
    review.freeze_panes = "A2"
    review.auto_filter.ref = review.dimensions
    decision_validation = DataValidation(
        type="list",
        formula1='"keep_gold,remove_relation,correct_entity,correct_relation,insufficient_context"',
        allow_blank=True,
    )
    review.add_data_validation(decision_validation)
    decision_validation.add(f"R2:S{max(2, review.max_row)}")
    for row in review.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    output_path = OUTPUT_DIR / "dev_gold_priority_reannotation.xlsx"
    workbook.save(output_path)
    print(f"rows={len(rows)}")
    print(output_path)


if __name__ == "__main__":
    main()
