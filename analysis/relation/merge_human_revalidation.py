from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[2]
EXPERIMENT_DIR = ROOT / "paper_revision" / "relation_improvement"
REVIEW_DIR = EXPERIMENT_DIR / "human_revalidation"
FILLED_DIR = REVIEW_DIR / "filled"

EXPERT1_PATH = FILLED_DIR / "01_RelationConfirmatory_Expert1_Blind_FILLED.xlsx"
EXPERT2_PATH = FILLED_DIR / "02_RelationConfirmatory_Expert2_Blind_FILLED.xlsx"

OUT_SUMMARY_JSON = REVIEW_DIR / "human_revalidation_pairwise_summary.json"
OUT_SUMMARY_CSV = REVIEW_DIR / "human_revalidation_pairwise_summary.csv"
OUT_CONFUSION_CSV = REVIEW_DIR / "human_revalidation_pairwise_confusion_matrix.csv"
OUT_MERGED_CSV = REVIEW_DIR / "human_revalidation_pairwise_merged.csv"
OUT_ADJUDICATION_XLSX = REVIEW_DIR / "03_RelationConfirmatory_Expert3_Disagreement_Adjudication.xlsx"

SHEET_NAME = "01_Blind_Ratings"
ALLOWED = ["present", "absent", "insufficient_context", "entity_or_type_error"]


def normalize(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def binary_label(label: str) -> str:
    return "present" if label == "present" else "not_present"


def read_blind_workbook(path: Path, expert_name: str) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"{path} does not contain sheet {SHEET_NAME!r}")
    sheet = workbook[SHEET_NAME]
    headers = [normalize(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    required = [
        "review_id",
        "sample_id",
        "full_passage_text",
        "local_evidence_window",
        "head_text",
        "head_type",
        "head_start",
        "head_end",
        "relation",
        "relation_definition",
        "tail_text",
        "tail_type",
        "tail_start",
        "tail_end",
        "relation_decision",
        "evidence_quote",
        "notes",
    ]
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError(f"{path} missing required columns: {missing}")
    idx = {name: headers.index(name) for name in required}

    rows: dict[str, dict[str, Any]] = {}
    invalid: list[tuple[str, str]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = {name: values[idx[name]] for name in required}
        rid = normalize(row["review_id"])
        if not rid:
            continue
        label = normalize(row["relation_decision"])
        if label not in ALLOWED:
            invalid.append((rid, label))
        if rid in rows:
            raise ValueError(f"{path} has duplicate review_id {rid}")
        row["relation_decision"] = label
        row["evidence_quote"] = normalize(row["evidence_quote"])
        row["notes"] = normalize(row["notes"])
        row["source_expert"] = expert_name
        rows[rid] = row
    workbook.close()
    if invalid:
        preview = invalid[:10]
        raise ValueError(f"{path} has invalid labels, first examples: {preview}")
    return rows


def cohen_kappa(labels1: list[str], labels2: list[str], categories: list[str]) -> float:
    if len(labels1) != len(labels2):
        raise ValueError("Label vectors must have the same length")
    n = len(labels1)
    if n == 0:
        return float("nan")
    observed = sum(a == b for a, b in zip(labels1, labels2)) / n
    c1 = Counter(labels1)
    c2 = Counter(labels2)
    expected = sum((c1[cat] / n) * (c2[cat] / n) for cat in categories)
    if expected == 1:
        return 1.0 if observed == 1 else float("nan")
    return (observed - expected) / (1 - expected)


def write_csv(path: Path, rows: list[dict[str, Any]], headers: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({header: row.get(header, "") for header in headers})


def style_sheet(sheet, max_row: int, decision_col: str | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {
        "A": 22,
        "B": 14,
        "C": 80,
        "D": 60,
        "E": 20,
        "F": 16,
        "G": 12,
        "H": 12,
        "I": 24,
        "J": 38,
        "K": 20,
        "L": 16,
        "M": 12,
        "N": 12,
        "O": 24,
        "P": 48,
        "Q": 42,
        "R": 24,
        "S": 48,
        "T": 42,
        "U": 20,
        "V": 24,
        "W": 60,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if decision_col is not None:
        validation = DataValidation(
            type="list",
            formula1='"present,absent,insufficient_context,entity_or_type_error"',
            allow_blank=True,
        )
        sheet.add_data_validation(validation)
        validation.add(f"{decision_col}2:{decision_col}{max(2, max_row)}")


def write_adjudication_workbook(
    disagreements: list[dict[str, Any]],
    agreements: list[dict[str, Any]],
    summary_rows: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "00_Instructions"
    instruction_rows = [
        ("Role", "Third expert adjudicator"),
        ("Task", "Only rows in 01_Disagreements need adjudication."),
        ("Blindness", "Do not inspect model predictions or previous gold labels."),
        ("Decision rule", "Choose the label best supported by the passage and entity boundaries."),
        ("present", "The passage explicitly supports the relation."),
        ("absent", "The passage does not support the relation."),
        ("insufficient_context", "The passage/OCR context is too damaged or incomplete to decide."),
        ("entity_or_type_error", "The entity span or entity type is invalid for this candidate."),
    ]
    for row in instruction_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 26
    instructions.column_dimensions["B"].width = 100
    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    summary = workbook.create_sheet("00_Summary")
    if summary_rows:
        headers = list(summary_rows[0].keys())
        summary.append(headers)
        for row in summary_rows:
            summary.append([row.get(header, "") for header in headers])
        style_sheet(summary, summary.max_row)

    disagreement_sheet = workbook.create_sheet("01_Disagreements")
    headers = [
        "review_id",
        "sample_id",
        "full_passage_text",
        "local_evidence_window",
        "head_text",
        "head_type",
        "head_start",
        "head_end",
        "relation",
        "relation_definition",
        "tail_text",
        "tail_type",
        "tail_start",
        "tail_end",
        "expert1_decision",
        "expert1_evidence",
        "expert1_notes",
        "expert2_decision",
        "expert2_evidence",
        "expert2_notes",
        "agreement_status",
        "adjudicated_decision",
        "adjudication_evidence_and_reason",
    ]
    disagreement_sheet.append(headers)
    for row in disagreements:
        disagreement_sheet.append([row.get(header, "") for header in headers])
    style_sheet(disagreement_sheet, disagreement_sheet.max_row, "V")

    agreement_sheet = workbook.create_sheet("02_Agreements_AutoAccepted")
    agreement_sheet.append(headers)
    for row in agreements:
        output = dict(row)
        output["adjudicated_decision"] = row["expert1_decision"]
        output["adjudication_evidence_and_reason"] = "Auto-accepted because Expert 1 and Expert 2 agreed."
        agreement_sheet.append([output.get(header, "") for header in headers])
    style_sheet(agreement_sheet, agreement_sheet.max_row)

    workbook.save(OUT_ADJUDICATION_XLSX)


def main() -> None:
    expert1 = read_blind_workbook(EXPERT1_PATH, "expert1")
    expert2 = read_blind_workbook(EXPERT2_PATH, "expert2")
    ids1 = set(expert1)
    ids2 = set(expert2)
    if ids1 != ids2:
        only1 = sorted(ids1 - ids2)[:10]
        only2 = sorted(ids2 - ids1)[:10]
        raise ValueError(f"review_id mismatch; only in expert1={only1}, only in expert2={only2}")

    merged: list[dict[str, Any]] = []
    agreements: list[dict[str, Any]] = []
    disagreements: list[dict[str, Any]] = []
    labels1: list[str] = []
    labels2: list[str] = []
    binary1: list[str] = []
    binary2: list[str] = []

    for rid in sorted(ids1):
        left = expert1[rid]
        right = expert2[rid]
        for key in [
            "sample_id",
            "head_text",
            "head_type",
            "head_start",
            "head_end",
            "relation",
            "tail_text",
            "tail_type",
            "tail_start",
            "tail_end",
        ]:
            if normalize(left[key]) != normalize(right[key]):
                raise ValueError(f"Candidate mismatch for {rid} in {key}")
        row = {
            "review_id": rid,
            "sample_id": left["sample_id"],
            "full_passage_text": left["full_passage_text"],
            "local_evidence_window": left["local_evidence_window"],
            "head_text": left["head_text"],
            "head_type": left["head_type"],
            "head_start": left["head_start"],
            "head_end": left["head_end"],
            "relation": left["relation"],
            "relation_definition": left["relation_definition"],
            "tail_text": left["tail_text"],
            "tail_type": left["tail_type"],
            "tail_start": left["tail_start"],
            "tail_end": left["tail_end"],
            "expert1_decision": left["relation_decision"],
            "expert1_evidence": left["evidence_quote"],
            "expert1_notes": left["notes"],
            "expert2_decision": right["relation_decision"],
            "expert2_evidence": right["evidence_quote"],
            "expert2_notes": right["notes"],
            "agreement_status": "agreement"
            if left["relation_decision"] == right["relation_decision"]
            else "disagreement",
            "adjudicated_decision": "",
            "adjudication_evidence_and_reason": "",
        }
        labels1.append(left["relation_decision"])
        labels2.append(right["relation_decision"])
        binary1.append(binary_label(left["relation_decision"]))
        binary2.append(binary_label(right["relation_decision"]))
        merged.append(row)
        if row["agreement_status"] == "agreement":
            agreements.append(row)
        else:
            disagreements.append(row)

    total = len(merged)
    exact_agreement = len(agreements) / total
    binary_agreement = sum(a == b for a, b in zip(binary1, binary2)) / total
    four_class_kappa = cohen_kappa(labels1, labels2, ALLOWED)
    binary_kappa = cohen_kappa(binary1, binary2, ["present", "not_present"])

    e1_counts = Counter(labels1)
    e2_counts = Counter(labels2)
    confusion_rows: list[dict[str, Any]] = []
    for a in ALLOWED:
        row = {"expert1_decision": a}
        for b in ALLOWED:
            row[b] = sum(x == a and y == b for x, y in zip(labels1, labels2))
        confusion_rows.append(row)

    summary = {
        "total_candidates": total,
        "expert1_nonblank": len(labels1),
        "expert2_nonblank": len(labels2),
        "exact_agreements": len(agreements),
        "exact_disagreements": len(disagreements),
        "exact_agreement_rate": exact_agreement,
        "four_class_cohen_kappa": four_class_kappa,
        "binary_agreement_rate_present_vs_not_present": binary_agreement,
        "binary_cohen_kappa_present_vs_not_present": binary_kappa,
        "expert1_present": e1_counts["present"],
        "expert1_absent": e1_counts["absent"],
        "expert1_insufficient_context": e1_counts["insufficient_context"],
        "expert1_entity_or_type_error": e1_counts["entity_or_type_error"],
        "expert2_present": e2_counts["present"],
        "expert2_absent": e2_counts["absent"],
        "expert2_insufficient_context": e2_counts["insufficient_context"],
        "expert2_entity_or_type_error": e2_counts["entity_or_type_error"],
    }

    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    OUT_SUMMARY_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(OUT_SUMMARY_CSV, [summary], list(summary.keys()))
    write_csv(OUT_CONFUSION_CSV, confusion_rows, ["expert1_decision"] + ALLOWED)
    merged_headers = list(merged[0].keys())
    write_csv(OUT_MERGED_CSV, merged, merged_headers)

    summary_rows = [{"metric": key, "value": value} for key, value in summary.items()]
    write_adjudication_workbook(disagreements, agreements, summary_rows)

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"wrote {OUT_ADJUDICATION_XLSX}")


if __name__ == "__main__":
    main()
