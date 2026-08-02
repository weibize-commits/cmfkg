from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import random
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
EXPERIMENT_DIR = ROOT / "paper_revision" / "relation_improvement"
DATA_DIR = EXPERIMENT_DIR / "data"
OUTPUT_DIR = EXPERIMENT_DIR / "outputs"
CONFIG_DIR = EXPERIMENT_DIR / "configs"
REVIEW_DIR = EXPERIMENT_DIR / "human_revalidation"
FINAL_DIR = REVIEW_DIR / "final"
FILLED_DIR = REVIEW_DIR / "filled"
PREDICTION_DIR = FINAL_DIR / "predictions"
BUILDER_PATH = ROOT / "paper_revision" / "tools" / "build_relation_baseline_package.py"

ADJUDICATION_PATH = (
    FILLED_DIR / "03_RelationConfirmatory_Expert3_Disagreement_Adjudication_FILLED.xlsx"
)
MERGED_CSV_PATH = REVIEW_DIR / "human_revalidation_pairwise_merged.csv"

SEED = 20260731
BOOTSTRAP_ROUNDS = 5000
ALLOWED_LABELS = {"present", "absent", "insufficient_context", "entity_or_type_error"}


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot import {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def write_jsonl(path: Path, records: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def clean_compact(value: Any) -> str:
    return "".join(str(value or "").split())


def clean_spaced(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def candidate_key(record: dict[str, Any]) -> tuple[Any, ...]:
    head = record["head"]
    tail = record["tail"]
    return (
        str(record["sample_id"]),
        int(head["start_char"]),
        int(head["end_char"]),
        clean_compact(head["text"]),
        str(head["type"]),
        str(record["relation"]),
        int(tail["start_char"]),
        int(tail["end_char"]),
        clean_compact(tail["text"]),
        str(tail["type"]),
    )


def review_id(record: dict[str, Any]) -> str:
    payload = json.dumps(candidate_key(record), ensure_ascii=False, separators=(",", ":"))
    return "REL-" + hashlib.sha256(payload.encode("utf-8")).hexdigest().upper()[:16]


def llm_digest(record: dict[str, Any]) -> str:
    head = record["head"]
    tail = record["tail"]
    key = (
        str(record["sample_id"]),
        int(head["start_char"]),
        int(head["end_char"]),
        clean_spaced(head["text"]),
        str(head["type"]),
        str(record["relation"]),
        int(tail["start_char"]),
        int(tail["end_char"]),
        clean_spaced(tail["text"]),
        str(tail["type"]),
    )
    payload = json.dumps(key, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest().upper()[:20]


def llm_support_score(decision: dict[str, Any]) -> float:
    label = str(decision.get("label") or "unsupported")
    confidence = float(decision.get("confidence") or 0.0)
    if label == "supported":
        return max(0.75, confidence)
    if label == "plausible":
        return 0.55 + 0.25 * confidence
    return 0.15 * (1.0 - confidence)


def prf_counts(tp: int, fp: int, fn: int) -> tuple[float, float, float]:
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = (
        2 * precision * recall / (precision + recall)
        if precision + recall
        else 0.0
    )
    return precision, recall, f1


def by_sample(records: Iterable[dict[str, Any]]) -> dict[str, set[tuple[Any, ...]]]:
    result: dict[str, set[tuple[Any, ...]]] = defaultdict(set)
    for record in records:
        result[str(record["sample_id"])].add(candidate_key(record))
    return result


def count_for_samples(
    gold_by_sample: dict[str, set[tuple[Any, ...]]],
    pred_by_sample: dict[str, set[tuple[Any, ...]]],
    sampled_ids: list[str],
) -> tuple[int, int, int]:
    tp = fp = fn = 0
    for sid in sampled_ids:
        gold = gold_by_sample.get(sid, set())
        pred = pred_by_sample.get(sid, set())
        tp += len(gold & pred)
        fp += len(pred - gold)
        fn += len(gold - pred)
    return tp, fp, fn


def bootstrap_metrics(
    gold_by_sample: dict[str, set[tuple[Any, ...]]],
    pred_by_sample: dict[str, set[tuple[Any, ...]]],
    sample_ids: list[str],
    seed: int,
) -> dict[str, tuple[float, float]]:
    rng = random.Random(seed)
    values = {"precision": [], "recall": [], "f1": []}
    for _ in range(BOOTSTRAP_ROUNDS):
        sampled = [rng.choice(sample_ids) for _ in sample_ids]
        metrics = prf_counts(*count_for_samples(gold_by_sample, pred_by_sample, sampled))
        for name, value in zip(["precision", "recall", "f1"], metrics):
            values[name].append(value)
    return {
        name: (
            float(np.percentile(items, 2.5)),
            float(np.percentile(items, 97.5)),
        )
        for name, items in values.items()
    }


def paired_bootstrap_difference(
    gold_by_sample: dict[str, set[tuple[Any, ...]]],
    left: dict[str, set[tuple[Any, ...]]],
    right: dict[str, set[tuple[Any, ...]]],
    sample_ids: list[str],
    seed: int,
) -> tuple[float, float, float, float]:
    observed_left = prf_counts(*count_for_samples(gold_by_sample, left, sample_ids))[2]
    observed_right = prf_counts(*count_for_samples(gold_by_sample, right, sample_ids))[2]
    rng = random.Random(seed)
    differences: list[float] = []
    for _ in range(BOOTSTRAP_ROUNDS):
        sampled = [rng.choice(sample_ids) for _ in sample_ids]
        left_f1 = prf_counts(*count_for_samples(gold_by_sample, left, sampled))[2]
        right_f1 = prf_counts(*count_for_samples(gold_by_sample, right, sampled))[2]
        differences.append(left_f1 - right_f1)
    low, high = np.percentile(differences, [2.5, 97.5])
    difference_array = np.array(differences)
    lower_tail = (int(np.sum(difference_array <= 0)) + 1) / (BOOTSTRAP_ROUNDS + 1)
    upper_tail = (int(np.sum(difference_array >= 0)) + 1) / (BOOTSTRAP_ROUNDS + 1)
    two_sided_p = 2 * min(lower_tail, upper_tail)
    return observed_left - observed_right, float(low), float(high), float(min(1.0, two_sided_p))


def read_csv_dicts(path: Path) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def read_disagreement_adjudications(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "01_Disagreements" not in workbook.sheetnames:
        raise RuntimeError("Adjudication workbook missing 01_Disagreements")
    sheet = workbook["01_Disagreements"]
    headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(max_row=1))]
    idx = {name: headers.index(name) for name in headers}
    required = [
        "review_id",
        "adjudicated_decision",
        "adjudication_evidence_and_reason",
    ]
    missing = [name for name in required if name not in idx]
    if missing:
        raise RuntimeError(f"Adjudication workbook missing columns: {missing}")
    rows: dict[str, dict[str, Any]] = {}
    invalid: list[tuple[str, str]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        rid = str(values[idx["review_id"]] or "").strip()
        if not rid:
            continue
        label = str(values[idx["adjudicated_decision"]] or "").strip()
        reason = str(values[idx["adjudication_evidence_and_reason"]] or "").strip()
        if label not in ALLOWED_LABELS:
            invalid.append((rid, label))
        rows[rid] = {
            "adjudicated_decision": label,
            "adjudication_evidence_and_reason": reason,
        }
    workbook.close()
    if invalid:
        raise RuntimeError(f"Invalid adjudicated decisions: {invalid[:10]}")
    return rows


def build_final_labels(candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    pairwise_rows = read_csv_dicts(MERGED_CSV_PATH)
    adjudications = read_disagreement_adjudications(ADJUDICATION_PATH)
    candidate_by_review_id = {review_id(candidate): candidate for candidate in candidates}
    if len(candidate_by_review_id) != len(candidates):
        raise RuntimeError("Duplicate review_id generated from candidates")

    final_rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in pairwise_rows:
        rid = str(row["review_id"]).strip()
        if rid not in candidate_by_review_id:
            raise RuntimeError(f"Unknown review_id in pairwise merge: {rid}")
        agreement_status = str(row["agreement_status"]).strip()
        expert1 = str(row["expert1_decision"]).strip()
        expert2 = str(row["expert2_decision"]).strip()
        if agreement_status == "agreement":
            if expert1 != expert2:
                raise RuntimeError(f"Agreement row has different expert decisions: {rid}")
            final_decision = expert1
            final_reason = "Expert 1 and Expert 2 agreed; auto-accepted."
            final_source = "two_expert_agreement"
        elif agreement_status == "disagreement":
            if rid not in adjudications:
                raise RuntimeError(f"Missing adjudication for disagreement: {rid}")
            final_decision = adjudications[rid]["adjudicated_decision"]
            final_reason = adjudications[rid]["adjudication_evidence_and_reason"]
            final_source = "third_expert_adjudication"
        else:
            raise RuntimeError(f"Unknown agreement_status {agreement_status!r} for {rid}")
        if final_decision not in ALLOWED_LABELS:
            raise RuntimeError(f"Invalid final decision {final_decision!r} for {rid}")

        candidate = candidate_by_review_id[rid]
        final_rows.append(
            {
                "review_id": rid,
                "sample_id": row["sample_id"],
                "head_text": row["head_text"],
                "head_type": row["head_type"],
                "head_start": row["head_start"],
                "head_end": row["head_end"],
                "relation": row["relation"],
                "tail_text": row["tail_text"],
                "tail_type": row["tail_type"],
                "tail_start": row["tail_start"],
                "tail_end": row["tail_end"],
                "expert1_decision": expert1,
                "expert1_evidence": row.get("expert1_evidence", ""),
                "expert1_notes": row.get("expert1_notes", ""),
                "expert2_decision": expert2,
                "expert2_evidence": row.get("expert2_evidence", ""),
                "expert2_notes": row.get("expert2_notes", ""),
                "agreement_status": agreement_status,
                "final_decision": final_decision,
                "final_source": final_source,
                "final_evidence_or_reason": final_reason,
                "candidate_key_json": json.dumps(
                    candidate_key(candidate),
                    ensure_ascii=False,
                    separators=(",", ":"),
                ),
            }
        )
        seen.add(rid)
    if seen != set(candidate_by_review_id):
        missing = sorted(set(candidate_by_review_id) - seen)[:10]
        raise RuntimeError(f"Final labels do not cover all candidates; missing {missing}")
    return sorted(final_rows, key=lambda value: value["review_id"])


def write_csv(path: Path, rows: list[dict[str, Any]], headers: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({header: row.get(header, "") for header in headers})


def filter_records(
    records: Iterable[dict[str, Any]],
    allowed_keys: set[tuple[Any, ...]] | None,
) -> list[dict[str, Any]]:
    if allowed_keys is None:
        return list(records)
    return [record for record in records if candidate_key(record) in allowed_keys]


def evaluate_methods(
    label_name: str,
    candidates: list[dict[str, Any]],
    sample_ids: list[str],
    method_records_all: dict[str, list[dict[str, Any]]],
    final_positive_keys: set[tuple[Any, ...]],
    allowed_eval_keys: set[tuple[Any, ...]] | None,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    eval_candidates = filter_records(candidates, allowed_eval_keys)
    eval_keys = {candidate_key(record) for record in eval_candidates}
    gold_records = [
        record
        for record in candidates
        if candidate_key(record) in final_positive_keys and candidate_key(record) in eval_keys
    ]
    gold_by_sample = by_sample(gold_records)

    pred_by_method: dict[str, dict[str, set[tuple[Any, ...]]]] = {}
    metric_rows: list[dict[str, Any]] = []
    for index, (method, records) in enumerate(method_records_all.items()):
        method_records = filter_records(records, eval_keys)
        pred_by_sample = by_sample(method_records)
        pred_by_method[method] = pred_by_sample
        tp, fp, fn = count_for_samples(gold_by_sample, pred_by_sample, sample_ids)
        precision, recall, f1 = prf_counts(tp, fp, fn)
        intervals = bootstrap_metrics(
            gold_by_sample,
            pred_by_sample,
            sample_ids,
            SEED + index,
        )
        metric_rows.append(
            {
                "analysis_set": label_name,
                "method": method,
                "n_passages": len(sample_ids),
                "gold_relations": len(gold_records),
                "candidate_pairs": len(eval_candidates),
                "predicted_relations": len(method_records),
                "tp": tp,
                "fp": fp,
                "fn": fn,
                "precision": precision,
                "precision_ci95_low": intervals["precision"][0],
                "precision_ci95_high": intervals["precision"][1],
                "recall": recall,
                "recall_ci95_low": intervals["recall"][0],
                "recall_ci95_high": intervals["recall"][1],
                "f1": f1,
                "f1_ci95_low": intervals["f1"][0],
                "f1_ci95_high": intervals["f1"][1],
            }
        )

    metrics = pd.DataFrame(metric_rows).sort_values("f1", ascending=False)

    relation_types = sorted({str(row["relation"]) for row in eval_candidates})
    by_relation_rows: list[dict[str, Any]] = []
    gold_keys = {candidate_key(row) for row in gold_records}
    for method, records in method_records_all.items():
        pred_keys = {candidate_key(row) for row in filter_records(records, eval_keys)}
        for relation in relation_types:
            relation_gold = {key for key in gold_keys if key[5] == relation}
            relation_pred = {key for key in pred_keys if key[5] == relation}
            tp = len(relation_gold & relation_pred)
            fp = len(relation_pred - relation_gold)
            fn = len(relation_gold - relation_pred)
            precision, recall, f1 = prf_counts(tp, fp, fn)
            by_relation_rows.append(
                {
                    "analysis_set": label_name,
                    "method": method,
                    "relation": relation,
                    "gold": len(relation_gold),
                    "predicted": len(relation_pred),
                    "tp": tp,
                    "fp": fp,
                    "fn": fn,
                    "precision": precision,
                    "recall": recall,
                    "f1": f1,
                }
            )

    paired_rows: list[dict[str, Any]] = []
    for comparison_index, right_name in enumerate(
        [
            "fixed_charngram_comparator",
            "local_symbolic_text",
            "distance_trigger_heuristic",
            "deepseek_supported_or_plausible",
        ]
    ):
        difference, low, high, p_value = paired_bootstrap_difference(
            gold_by_sample,
            pred_by_method["cmfkg_frozen_hybrid"],
            pred_by_method[right_name],
            sample_ids,
            SEED + 100 + comparison_index,
        )
        paired_rows.append(
            {
                "analysis_set": label_name,
                "left_method": "cmfkg_frozen_hybrid",
                "right_method": right_name,
                "f1_difference": difference,
                "ci95_low": low,
                "ci95_high": high,
                "paired_bootstrap_two_sided_p": p_value,
                "bootstrap_unit": "passage",
                "bootstrap_rounds": BOOTSTRAP_ROUNDS,
            }
        )
    return metrics, pd.DataFrame(paired_rows), pd.DataFrame(by_relation_rows)


def build_method_records(candidates: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    builder = load_module("relation_builder_human_revalidated", BUILDER_PATH)
    local_freeze = json.loads(
        (CONFIG_DIR / "local_model_freeze.json").read_text(encoding="utf-8")
    )
    comparator_freeze = json.loads(
        (CONFIG_DIR / "fixed_comparator_freeze.json").read_text(encoding="utf-8")
    )
    policy_freeze = json.loads(
        (CONFIG_DIR / "llm_policy_freeze.json").read_text(encoding="utf-8")
    )
    passages = read_jsonl(DATA_DIR / "test_passages.jsonl")
    local_scores_raw = read_jsonl(OUTPUT_DIR / "confirmatory_test_scores_blinded.jsonl")
    comparator_scores_raw = read_jsonl(OUTPUT_DIR / "fixed_charngram_test_scores_blinded.jsonl")
    decisions = read_jsonl(OUTPUT_DIR / "deepseek_triage_test_decisions.jsonl")
    texts = {
        str(row["sample_id"]): {"frozen_text": str(row["frozen_text"])}
        for row in passages
    }

    local_score_by_key = {candidate_key(row): float(row["score"]) for row in local_scores_raw}
    comparator_score_by_key = {
        candidate_key(row): float(row["score"]) for row in comparator_scores_raw
    }
    decision_by_digest = {str(row["candidate_digest"]): row for row in decisions}
    if (
        len(local_score_by_key) != len(candidates)
        or len(comparator_score_by_key) != len(candidates)
        or len(decision_by_digest) != len(candidates)
    ):
        raise RuntimeError("Candidate, score, or LLM-decision cardinality mismatch")

    method_records: dict[str, list[dict[str, Any]]] = {
        "typed_cooccurrence": candidates[:],
        "distance_trigger_heuristic": [],
        "fixed_charngram_comparator": [],
        "local_symbolic_text": [],
        "deepseek_supported_only": [],
        "deepseek_supported_or_plausible": [],
        "cmfkg_frozen_hybrid": [],
    }
    local_thresholds = {
        str(key): float(value)
        for key, value in local_freeze["relation_thresholds"].items()
    }
    policy_family = str(policy_freeze["selected_family"])
    policy = policy_freeze["selected_parameters"]

    for candidate in candidates:
        key = candidate_key(candidate)
        local_score = local_score_by_key[key]
        comparator_score = comparator_score_by_key[key]
        heuristic_keep, heuristic_score, _ = builder.heuristic_score(candidate, texts)
        decision = decision_by_digest[llm_digest(candidate)]
        llm_score = llm_support_score(decision)
        label = str(decision["label"])

        if heuristic_keep:
            method_records["distance_trigger_heuristic"].append(candidate)
        if comparator_score >= float(comparator_freeze["threshold_selected_on_dev"]):
            method_records["fixed_charngram_comparator"].append(candidate)
        if local_score >= local_thresholds.get(
            str(candidate["relation"]),
            float(local_freeze["selected_local_model"]["global_threshold"]),
        ):
            method_records["local_symbolic_text"].append(candidate)
        if label == "supported":
            method_records["deepseek_supported_only"].append(candidate)
        if label in {"supported", "plausible"}:
            method_records["deepseek_supported_or_plausible"].append(candidate)

        if policy_family == "local_only":
            hybrid_positive = local_score >= float(policy["threshold"])
        elif policy_family == "local_relation_thresholds":
            hybrid_positive = local_score >= local_thresholds.get(
                str(candidate["relation"]),
                float(local_freeze["selected_local_model"]["global_threshold"]),
            )
        elif policy_family == "relation_threshold_llm_rescue_veto":
            relation_threshold = local_thresholds.get(
                str(candidate["relation"]),
                float(local_freeze["selected_local_model"]["global_threshold"]),
            )
            base_positive = local_score >= relation_threshold
            llm_accept = label == "supported" or (
                bool(policy["accept_plausible"]) and label == "plausible"
            )
            rescued = (
                (not base_positive)
                and llm_accept
                and local_score >= max(0.0, relation_threshold - float(policy["rescue_margin"]))
            )
            veto_margin = policy.get("veto_margin")
            vetoed = (
                base_positive
                and veto_margin is not None
                and label == "unsupported"
                and local_score < relation_threshold + float(veto_margin)
            )
            hybrid_positive = (base_positive or rescued) and not vetoed
        elif policy_family == "weighted_score":
            combined = (
                float(policy["local_weight"]) * local_score
                + float(policy["llm_weight"]) * llm_score
                + float(policy["heuristic_weight"]) * float(heuristic_score)
            )
            hybrid_positive = combined >= float(policy["threshold"])
        elif policy_family == "protected_local_gate":
            llm_accept = label == "supported" or (
                bool(policy["accept_plausible"]) and label == "plausible"
            )
            hybrid_positive = local_score >= float(policy["upper"]) or (
                local_score >= float(policy["lower"]) and llm_accept
            )
        else:
            raise RuntimeError(f"Unknown frozen policy family: {policy_family}")
        if hybrid_positive:
            method_records["cmfkg_frozen_hybrid"].append(candidate)
    return method_records


def main() -> None:
    FINAL_DIR.mkdir(parents=True, exist_ok=True)
    PREDICTION_DIR.mkdir(parents=True, exist_ok=True)

    passages = read_jsonl(DATA_DIR / "test_passages.jsonl")
    candidates = read_jsonl(DATA_DIR / "test_candidates.jsonl")
    sample_ids = [str(row["sample_id"]) for row in passages]
    candidate_by_review_id = {review_id(candidate): candidate for candidate in candidates}
    final_rows = build_final_labels(candidates)
    final_by_review_id = {row["review_id"]: row for row in final_rows}
    if set(final_by_review_id) != set(candidate_by_review_id):
        raise RuntimeError("Final labels and candidate review IDs do not match")

    final_label_headers = list(final_rows[0].keys())
    write_csv(FINAL_DIR / "human_revalidated_final_labels.csv", final_rows, final_label_headers)

    final_positive_records = [
        candidate_by_review_id[row["review_id"]]
        for row in final_rows
        if row["final_decision"] == "present"
    ]
    write_jsonl(
        FINAL_DIR / "human_revalidated_gold_positive.jsonl",
        final_positive_records,
    )

    method_records_all = build_method_records(candidates)
    for method, records in method_records_all.items():
        write_jsonl(PREDICTION_DIR / f"{method}.jsonl", records)

    final_positive_keys = {candidate_key(record) for record in final_positive_records}
    decidable_keys = {
        candidate_key(candidate_by_review_id[row["review_id"]])
        for row in final_rows
        if row["final_decision"] in {"present", "absent"}
    }

    primary_metrics, primary_paired, primary_by_relation = evaluate_methods(
        "primary_present_vs_all_nonpresent",
        candidates,
        sample_ids,
        method_records_all,
        final_positive_keys,
        allowed_eval_keys=None,
    )
    valid_metrics, valid_paired, valid_by_relation = evaluate_methods(
        "sensitivity_present_vs_absent_only",
        candidates,
        sample_ids,
        method_records_all,
        final_positive_keys,
        allowed_eval_keys=decidable_keys,
    )

    all_metrics = pd.concat([primary_metrics, valid_metrics], ignore_index=True)
    all_paired = pd.concat([primary_paired, valid_paired], ignore_index=True)
    all_by_relation = pd.concat(
        [primary_by_relation, valid_by_relation],
        ignore_index=True,
    )
    all_metrics.to_csv(
        FINAL_DIR / "human_revalidated_relation_metrics.csv",
        index=False,
        encoding="utf-8-sig",
    )
    all_paired.to_csv(
        FINAL_DIR / "human_revalidated_paired_bootstrap.csv",
        index=False,
        encoding="utf-8-sig",
    )
    all_by_relation.to_csv(
        FINAL_DIR / "human_revalidated_by_relation.csv",
        index=False,
        encoding="utf-8-sig",
    )

    label_counts = Counter(row["final_decision"] for row in final_rows)
    source_counts = Counter(row["final_source"] for row in final_rows)
    summary_rows = [
        {"metric": "n_passages", "value": len(sample_ids)},
        {"metric": "candidate_pairs", "value": len(candidates)},
        {"metric": "final_present", "value": label_counts["present"]},
        {"metric": "final_absent", "value": label_counts["absent"]},
        {"metric": "final_insufficient_context", "value": label_counts["insufficient_context"]},
        {"metric": "final_entity_or_type_error", "value": label_counts["entity_or_type_error"]},
        {"metric": "two_expert_agreement_auto_accepted", "value": source_counts["two_expert_agreement"]},
        {"metric": "third_expert_adjudicated", "value": source_counts["third_expert_adjudication"]},
    ]
    summary_table = pd.DataFrame(summary_rows)
    with pd.ExcelWriter(
        FINAL_DIR / "human_revalidated_relation_experiment.xlsx",
        engine="openpyxl",
    ) as writer:
        summary_table.to_excel(writer, sheet_name="00_summary", index=False)
        all_metrics.to_excel(writer, sheet_name="01_metrics", index=False)
        all_paired.to_excel(writer, sheet_name="02_paired_bootstrap", index=False)
        all_by_relation.to_excel(writer, sheet_name="03_by_relation", index=False)
        pd.DataFrame(final_rows).to_excel(writer, sheet_name="04_final_labels", index=False)

    primary = primary_metrics.set_index("method")
    valid = valid_metrics.set_index("method")
    primary_hybrid = primary.loc["cmfkg_frozen_hybrid"]
    primary_comparator = primary.loc["fixed_charngram_comparator"]
    valid_hybrid = valid.loc["cmfkg_frozen_hybrid"]
    valid_comparator = valid.loc["fixed_charngram_comparator"]
    summary_md = f"""# Human-revalidated relation benchmark

- Human-revalidated confirmatory set: {len(sample_ids)} passages and {len(candidates)} candidate relation pairs.
- Final labels: present={label_counts['present']}, absent={label_counts['absent']}, insufficient_context={label_counts['insufficient_context']}, entity_or_type_error={label_counts['entity_or_type_error']}.
- Adjudication: {source_counts['two_expert_agreement']} labels were auto-accepted after two-expert agreement; {source_counts['third_expert_adjudication']} disagreements were resolved by the third expert.
- Primary analysis: present versus all non-present labels; bootstrap unit=passage; resamples={BOOTSTRAP_ROUNDS}.
- Primary CMFKG hybrid: precision={primary_hybrid['precision']:.4f}, recall={primary_hybrid['recall']:.4f}, F1={primary_hybrid['f1']:.4f} (95% CI {primary_hybrid['f1_ci95_low']:.4f}-{primary_hybrid['f1_ci95_high']:.4f}).
- Primary fixed char-ngram comparator: precision={primary_comparator['precision']:.4f}, recall={primary_comparator['recall']:.4f}, F1={primary_comparator['f1']:.4f} (95% CI {primary_comparator['f1_ci95_low']:.4f}-{primary_comparator['f1_ci95_high']:.4f}).
- Sensitivity analysis excluding insufficient_context and entity_or_type_error: CMFKG F1={valid_hybrid['f1']:.4f}; fixed comparator F1={valid_comparator['f1']:.4f}.
"""
    (FINAL_DIR / "human_revalidated_summary.md").write_text(summary_md, encoding="utf-8")

    manifest = {
        "status": "human_revalidated_gold_evaluated",
        "evaluated_at_utc": datetime.now(timezone.utc).isoformat(),
        "independent_unit": "passage",
        "bootstrap_rounds": BOOTSTRAP_ROUNDS,
        "seed": SEED,
        "test_passages": len(sample_ids),
        "candidate_pairs": len(candidates),
        "label_counts": dict(label_counts),
        "source_counts": dict(source_counts),
        "expert3_adjudication_sha256": sha256(ADJUDICATION_PATH),
        "pairwise_merged_sha256": sha256(MERGED_CSV_PATH),
        "final_labels_sha256": sha256(FINAL_DIR / "human_revalidated_final_labels.csv"),
        "gold_positive_sha256": sha256(FINAL_DIR / "human_revalidated_gold_positive.jsonl"),
        "metrics_sha256": sha256(FINAL_DIR / "human_revalidated_relation_metrics.csv"),
        "paired_bootstrap_sha256": sha256(
            FINAL_DIR / "human_revalidated_paired_bootstrap.csv"
        ),
    }
    (FINAL_DIR / "human_revalidated_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(summary_md)
    print(FINAL_DIR)


if __name__ == "__main__":
    main()
