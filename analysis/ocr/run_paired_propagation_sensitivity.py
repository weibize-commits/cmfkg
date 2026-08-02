from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import math
import os
import pickle
import random
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import torch
import torch.nn as nn
from openpyxl import load_workbook
from torch.utils.data import DataLoader, Dataset
from torchcrf import CRF
from transformers import AutoModel, AutoTokenizer


os.environ.setdefault("HF_HUB_DISABLE_TELEMETRY", "1")
os.environ.setdefault("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")

ROOT = Path(__file__).resolve().parents[2]
OCR_ROOT = ROOT / "paper_revision" / "ocr_quantitative"
INPUT_XLSX = OCR_ROOT / "inputs" / "ocr_manual_transcription_filled.xlsx"
OUTPUT_DIR = OCR_ROOT / "outputs" / "paired_propagation"
BUILDER_PATH = ROOT / "paper_revision" / "tools" / "build_relation_baseline_package.py"
RELATION_MODEL_PATH = (
    ROOT
    / "paper_revision"
    / "relation_improvement"
    / "models"
    / "fixed_charngram_comparator.pkl"
)
RELATION_FREEZE_PATH = (
    ROOT
    / "paper_revision"
    / "relation_improvement"
    / "configs"
    / "fixed_comparator_freeze.json"
)
LABELS_PATH = ROOT / "labels.txt"

MAX_LENGTH = 256
WINDOW_CHARS = 254
WINDOW_OVERLAP = 48
BATCH_SIZE = 12
BOOTSTRAP_REPS = 10_000
BOOTSTRAP_SEED = 20260801

MODELS = {
    "SikuRoBERTa-CRF_archived": {
        "checkpoint": ROOT / "models" / "sikuroberta_crf" / "best_model.pt",
        "base_model": "SIKU-BERT/sikuroberta",
    },
    "Chinese-RoBERTa-CRF_archived": {
        "checkpoint": ROOT / "models" / "chinese_roberta_crf" / "best_model.pt",
        "base_model": "hfl/chinese-roberta-wwm-ext",
    },
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def read_text(path: str | Path) -> str:
    return Path(path).read_text(encoding="utf-8", errors="replace")


def normalize_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    return re.sub(r"\s+", "", str(value).replace("\u3000", ""))


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot import {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def workbook_rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    header = [cell.value for cell in worksheet[1]]
    rows: list[dict[str, Any]] = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        if any(value not in (None, "") for value in values):
            rows.append({header[index]: values[index] for index in range(len(header))})
    return rows


def load_paired_pages() -> list[dict[str, str]]:
    pages: list[dict[str, str]] = []
    for row in workbook_rows(INPUT_XLSX, "ManualTranscription"):
        reference = normalize_text(row.get("adjudicated_transcription"))
        cleaned_path = str(row.get("cleaned_ocr_page_text_path") or "")
        raw_path = str(row.get("raw_ocr_page_text_path") or "")
        if not reference or not raw_path or not cleaned_path:
            continue
        if not Path(raw_path).exists() or not Path(cleaned_path).exists():
            continue
        raw = normalize_text(read_text(raw_path))
        cleaned = normalize_text(read_text(cleaned_path))
        if not raw or not cleaned:
            continue
        pages.append(
            {
                "sample_id": str(row["sample_id"]),
                "source_key": str(row.get("source_key") or ""),
                "raw": raw,
                "cleaned": cleaned,
                "reference": reference,
            }
        )
    if len(pages) != 17:
        raise RuntimeError(f"Expected 17 paired pages, found {len(pages)}")
    return pages


def load_labels() -> tuple[list[str], dict[str, int], dict[int, str]]:
    labels = [x.strip() for x in LABELS_PATH.read_text(encoding="utf-8").splitlines() if x.strip()]
    return labels, {label: index for index, label in enumerate(labels)}, {
        index: label for index, label in enumerate(labels)
    }


def text_windows(text: str) -> list[dict[str, Any]]:
    if not text:
        return []
    stride = WINDOW_CHARS - WINDOW_OVERLAP
    windows: list[dict[str, Any]] = []
    start = 0
    while start < len(text):
        end = min(start + WINDOW_CHARS, len(text))
        windows.append({"start": start, "end": end, "text": text[start:end]})
        if end >= len(text):
            break
        start += stride
    return windows


class WindowDataset(Dataset):
    def __init__(self, records, tokenizer):
        self.records = records
        self.tokenizer = tokenizer

    def __len__(self):
        return len(self.records)

    def __getitem__(self, index):
        record = self.records[index]
        chars = list(record["text"])
        tokens = ["[CLS]"] + chars + ["[SEP]"]
        unk_id = self.tokenizer.unk_token_id
        input_ids = [
            token_id if token_id is not None else unk_id
            for token_id in self.tokenizer.convert_tokens_to_ids(tokens)
        ]
        attention_mask = [1] * len(input_ids)
        pad_length = MAX_LENGTH - len(input_ids)
        input_ids += [self.tokenizer.pad_token_id] * pad_length
        attention_mask += [0] * pad_length
        return {
            "input_ids": torch.tensor(input_ids, dtype=torch.long),
            "attention_mask": torch.tensor(attention_mask, dtype=torch.long),
            "record_index": index,
        }


class BertCRF(nn.Module):
    def __init__(self, base_model: str, num_labels: int):
        super().__init__()
        self.bert = AutoModel.from_pretrained(base_model)
        self.dropout = nn.Dropout(0.1)
        self.classifier = nn.Linear(self.bert.config.hidden_size, num_labels)
        self.crf = CRF(num_labels, batch_first=True)

    def forward(self, input_ids, attention_mask):
        outputs = self.bert(input_ids=input_ids, attention_mask=attention_mask)
        emissions = self.classifier(self.dropout(outputs.last_hidden_state))
        return self.crf.decode(emissions, mask=attention_mask.bool())


def tag_entities(tags: list[str], text: str) -> list[dict[str, Any]]:
    entities: list[dict[str, Any]] = []
    start: int | None = None
    label: str | None = None
    for index, tag in enumerate(tags + ["O"]):
        if tag.startswith("B-"):
            if start is not None and label is not None:
                entities.append(
                    {
                        "start_char": start,
                        "end_char": index,
                        "text": text[start:index],
                        "type": label,
                    }
                )
            start = index
            label = tag[2:]
        elif tag.startswith("I-") and start is not None and tag[2:] == label:
            continue
        else:
            if start is not None and label is not None:
                entities.append(
                    {
                        "start_char": start,
                        "end_char": index,
                        "text": text[start:index],
                        "type": label,
                    }
                )
            start = None
            label = None
    return entities


def infer_entities(
    model_name: str,
    config: dict[str, Any],
    pages: list[dict[str, str]],
) -> dict[tuple[str, str], list[dict[str, Any]]]:
    labels, _, id2label = load_labels()
    tokenizer = AutoTokenizer.from_pretrained(config["base_model"])
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    model = BertCRF(config["base_model"], len(labels)).to(device)
    checkpoint = torch.load(config["checkpoint"], weights_only=False, map_location=device)
    model.load_state_dict(checkpoint["model_state_dict"])
    model.eval()

    records: list[dict[str, Any]] = []
    for page in pages:
        for variant in ["raw", "cleaned", "reference"]:
            windows = text_windows(page[variant])
            for window_index, window in enumerate(windows):
                records.append(
                    {
                        "sample_id": page["sample_id"],
                        "variant": variant,
                        "full_text": page[variant],
                        "window_index": window_index,
                        "window_count": len(windows),
                        **window,
                    }
                )

    dataset = WindowDataset(records, tokenizer)
    loader = DataLoader(dataset, batch_size=BATCH_SIZE, shuffle=False)
    by_page: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    with torch.no_grad():
        for batch in loader:
            input_ids = batch["input_ids"].to(device)
            attention_mask = batch["attention_mask"].to(device)
            predictions = model(input_ids, attention_mask)
            for row_index, prediction in zip(batch["record_index"].tolist(), predictions):
                record = records[row_index]
                tags = [id2label[tag_id] for tag_id in prediction[1:-1]]
                tags = tags[: len(record["text"])]
                local_entities = tag_entities(tags, record["text"])
                for entity in local_entities:
                    midpoint = (entity["start_char"] + entity["end_char"]) / 2
                    left_margin = 0 if record["window_index"] == 0 else WINDOW_OVERLAP / 2
                    right_margin = (
                        len(record["text"])
                        if record["window_index"] == record["window_count"] - 1
                        else len(record["text"]) - WINDOW_OVERLAP / 2
                    )
                    if not (left_margin <= midpoint < right_margin):
                        continue
                    global_entity = dict(entity)
                    global_entity["start_char"] += int(record["start"])
                    global_entity["end_char"] += int(record["start"])
                    by_page[(record["sample_id"], record["variant"])].append(global_entity)

    for key, values in list(by_page.items()):
        deduplicated = {
            (
                int(entity["start_char"]),
                int(entity["end_char"]),
                str(entity["text"]),
                str(entity["type"]),
            ): entity
            for entity in values
        }
        by_page[key] = sorted(
            deduplicated.values(),
            key=lambda entity: (
                int(entity["start_char"]),
                int(entity["end_char"]),
                str(entity["type"]),
            ),
        )
    print(f"Completed NER inference: {model_name}", flush=True)
    del model
    if torch.cuda.is_available():
        torch.cuda.empty_cache()
    return by_page


def entity_signature(entity: dict[str, Any]) -> tuple[str, str]:
    return normalize_text(entity["text"]), str(entity["type"])


def relation_signature(record: dict[str, Any]) -> tuple[str, str, str, str, str]:
    return (
        str(record["relation"]),
        normalize_text(record["head"]["text"]),
        str(record["head"]["type"]),
        normalize_text(record["tail"]["text"]),
        str(record["tail"]["type"]),
    )


def rule_signature(record: dict[str, Any]) -> tuple[str, str, str, str, str]:
    return (
        str(record["rule"]),
        normalize_text(record["head_text"]),
        str(record["head_type"]),
        normalize_text(record["tail_text"]),
        str(record["tail_type"]),
    )


def set_counts(left: set[Any], reference: set[Any]) -> tuple[int, int, int]:
    return len(left & reference), len(left - reference), len(reference - left)


def prf_from_counts(tp: int, fp: int, fn: int) -> tuple[float, float, float]:
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return precision, recall, f1


def generate_relations(
    builder,
    model,
    threshold: float,
    sample_id: str,
    text: str,
    entities: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    texts = {sample_id: {"frozen_text": text}}
    candidates = builder.generate_candidates([sample_id], texts, {sample_id: entities})
    if not candidates:
        return [], []
    features = [builder.feature_text(candidate, texts) for candidate in candidates]
    scores = model.predict_proba(features)[:, 1]
    accepted = [
        candidate
        for candidate, score in zip(candidates, scores)
        if float(score) >= threshold
    ]
    return candidates, accepted


def apply_ab_rules(relations: list[dict[str, Any]]) -> list[dict[str, Any]]:
    edges: dict[str, set[tuple[tuple[str, str], tuple[str, str]]]] = defaultdict(set)
    for record in relations:
        head = (normalize_text(record["head"]["text"]), str(record["head"]["type"]))
        tail = (normalize_text(record["tail"]["text"]), str(record["tail"]["type"]))
        edges[str(record["relation"])].add((head, tail))

    output: set[tuple[str, tuple[str, str], tuple[str, str]]] = set()
    for recipe, ingredient in edges["containsIngredient"]:
        for source_ingredient, nature in edges["hasNature"]:
            if ingredient == source_ingredient:
                output.add(("A1", recipe, nature))
        for source_ingredient, flavour in edges["hasFlavor"]:
            if ingredient == source_ingredient:
                output.add(("A2", recipe, flavour))
        for source_ingredient, meridian in edges["entersmeridian"]:
            if ingredient == source_ingredient:
                output.add(("A3", recipe, meridian))

    for ingredient, meridian in edges["entersmeridian"]:
        for source_meridian, organ in edges["correspondsToOrgan"]:
            if meridian != source_meridian:
                continue
            for disease, affected_organ in edges["affectsOrgan"]:
                if organ == affected_organ:
                    output.add(("B1", ingredient, disease))
    for recipe, ingredient in edges["containsIngredient"]:
        for source_ingredient, disease in edges["treats"]:
            if ingredient == source_ingredient:
                output.add(("B2", recipe, disease))
    for disease, recipe in edges["treatedByRecipe"]:
        for source_recipe, ingredient in edges["containsIngredient"]:
            if recipe == source_recipe:
                output.add(("B3", ingredient, disease))

    return [
        {
            "rule": rule,
            "head_text": head[0],
            "head_type": head[1],
            "tail_text": tail[0],
            "tail_type": tail[1],
        }
        for rule, head, tail in sorted(output)
    ]


def aggregate_metric(
    page_rows: list[dict[str, Any]],
    prefix: str,
    variant: str,
) -> tuple[float, float, float]:
    tp = sum(int(row[f"{prefix}_{variant}_tp"]) for row in page_rows)
    fp = sum(int(row[f"{prefix}_{variant}_fp"]) for row in page_rows)
    fn = sum(int(row[f"{prefix}_{variant}_fn"]) for row in page_rows)
    return prf_from_counts(tp, fp, fn)


def bootstrap_delta(
    page_rows: list[dict[str, Any]], prefix: str
) -> tuple[float, float, float]:
    rng = random.Random(BOOTSTRAP_SEED + sum(ord(ch) for ch in prefix))
    deltas: list[float] = []
    for _ in range(BOOTSTRAP_REPS):
        sample = [page_rows[rng.randrange(len(page_rows))] for _ in page_rows]
        raw_f1 = aggregate_metric(sample, prefix, "raw")[2]
        cleaned_f1 = aggregate_metric(sample, prefix, "cleaned")[2]
        deltas.append(cleaned_f1 - raw_f1)
    deltas.sort()
    lo = deltas[int(0.025 * BOOTSTRAP_REPS)]
    hi = deltas[min(int(0.975 * BOOTSTRAP_REPS), BOOTSTRAP_REPS - 1)]
    point = aggregate_metric(page_rows, prefix, "cleaned")[2] - aggregate_metric(
        page_rows, prefix, "raw"
    )[2]
    return point, lo, hi


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


def analyse_model(
    model_name: str,
    config: dict[str, Any],
    pages: list[dict[str, str]],
    builder,
    relation_model,
    relation_threshold: float,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    entities_by_page = infer_entities(model_name, config, pages)
    page_rows: list[dict[str, Any]] = []
    prediction_rows: list[dict[str, Any]] = []

    for page in pages:
        sid = page["sample_id"]
        variant_payload: dict[str, dict[str, Any]] = {}
        for variant in ["raw", "cleaned", "reference"]:
            entities = entities_by_page.get((sid, variant), [])
            candidates, relations = generate_relations(
                builder,
                relation_model,
                relation_threshold,
                f"{sid}:{variant}",
                page[variant],
                entities,
            )
            rules = apply_ab_rules(relations)
            variant_payload[variant] = {
                "entities": entities,
                "relation_candidates": candidates,
                "relations": relations,
                "rule_outputs": rules,
            }
            prediction_rows.append(
                {
                    "model": model_name,
                    "sample_id": sid,
                    "variant": variant,
                    "text_length": len(page[variant]),
                    **variant_payload[variant],
                }
            )

        row: dict[str, Any] = {
            "model": model_name,
            "sample_id": sid,
            "source_key": page["source_key"],
            "raw_chars": len(page["raw"]),
            "cleaned_chars": len(page["cleaned"]),
            "reference_chars": len(page["reference"]),
        }
        signature_functions = {
            "entity": ("entities", entity_signature),
            "relation": ("relations", relation_signature),
            "rule": ("rule_outputs", rule_signature),
        }
        for prefix, (payload_key, signature_function) in signature_functions.items():
            reference_set = {
                signature_function(item)
                for item in variant_payload["reference"][payload_key]
            }
            row[f"{prefix}_reference_count"] = len(reference_set)
            for variant in ["raw", "cleaned"]:
                variant_set = {
                    signature_function(item)
                    for item in variant_payload[variant][payload_key]
                }
                tp, fp, fn = set_counts(variant_set, reference_set)
                precision, recall, f1 = prf_from_counts(tp, fp, fn)
                row[f"{prefix}_{variant}_count"] = len(variant_set)
                row[f"{prefix}_{variant}_tp"] = tp
                row[f"{prefix}_{variant}_fp"] = fp
                row[f"{prefix}_{variant}_fn"] = fn
                row[f"{prefix}_{variant}_stability_precision"] = precision
                row[f"{prefix}_{variant}_stability_recall"] = recall
                row[f"{prefix}_{variant}_stability_f1"] = f1
        page_rows.append(row)

    summary_rows: list[dict[str, Any]] = []
    for prefix in ["entity", "relation", "rule"]:
        raw_precision, raw_recall, raw_f1 = aggregate_metric(page_rows, prefix, "raw")
        cleaned_precision, cleaned_recall, cleaned_f1 = aggregate_metric(
            page_rows, prefix, "cleaned"
        )
        delta, delta_lo, delta_hi = bootstrap_delta(page_rows, prefix)
        summary_rows.append(
            {
                "model": model_name,
                "output_level": prefix,
                "n_pages": len(page_rows),
                "raw_stability_precision": raw_precision,
                "raw_stability_recall": raw_recall,
                "raw_stability_f1": raw_f1,
                "cleaned_stability_precision": cleaned_precision,
                "cleaned_stability_recall": cleaned_recall,
                "cleaned_stability_f1": cleaned_f1,
                "cleaned_minus_raw_stability_f1": delta,
                "bootstrap_ci_low": delta_lo,
                "bootstrap_ci_high": delta_hi,
                "reference_definition": "same frozen model output on adjudicated manual transcription",
                "interpretation": "prediction stability only; not human-annotated accuracy",
            }
        )
    return page_rows, summary_rows, prediction_rows


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    pages = load_paired_pages()
    builder = load_module("ocr_relation_builder", BUILDER_PATH)
    with RELATION_MODEL_PATH.open("rb") as handle:
        relation_model = pickle.load(handle)
    relation_freeze = json.loads(RELATION_FREEZE_PATH.read_text(encoding="utf-8"))
    relation_threshold = float(relation_freeze["threshold_selected_on_dev"])

    all_page_rows: list[dict[str, Any]] = []
    all_summary_rows: list[dict[str, Any]] = []
    all_prediction_rows: list[dict[str, Any]] = []
    for model_name, config in MODELS.items():
        page_rows, summary_rows, prediction_rows = analyse_model(
            model_name,
            config,
            pages,
            builder,
            relation_model,
            relation_threshold,
        )
        all_page_rows.extend(page_rows)
        all_summary_rows.extend(summary_rows)
        all_prediction_rows.extend(prediction_rows)

    page_path = OUTPUT_DIR / "paired_propagation_page_results.csv"
    summary_path = OUTPUT_DIR / "paired_propagation_summary.csv"
    prediction_path = OUTPUT_DIR / "paired_propagation_predictions.jsonl"
    write_csv(page_path, all_page_rows)
    write_csv(summary_path, all_summary_rows)
    write_jsonl(prediction_path, all_prediction_rows)

    manifest = {
        "status": "complete_paired_propagation_sensitivity",
        "scope": "17 pages with raw OCR, cleaned OCR, and adjudicated manual transcription",
        "models": {
            name: {
                "base_model": config["base_model"],
                "checkpoint": str(config["checkpoint"]),
                "checkpoint_sha256": sha256(config["checkpoint"]),
            }
            for name, config in MODELS.items()
        },
        "windowing": {
            "maximum_model_length": MAX_LENGTH,
            "window_characters": WINDOW_CHARS,
            "overlap_characters": WINDOW_OVERLAP,
            "overlap_resolution": "retain an entity when its midpoint lies in the window's non-overlap ownership region",
        },
        "relation_model": {
            "path": str(RELATION_MODEL_PATH),
            "sha256": sha256(RELATION_MODEL_PATH),
            "threshold": relation_threshold,
        },
        "rule_scope": "audited A1-A3 and B1-B3 Horn-style candidates applied to frozen relation predictions",
        "important_limitations": [
            "The adjudicated transcription has no independent entity or relation annotations for these pages.",
            "Stability against the same model on adjudicated text is not NER, relation, or rule accuracy.",
            "The 17-page paired subset does not represent all 51 adjudicated pages or the full corpus.",
            "C-class OCR-cause precision is not estimated because the 341 archived flags have no item-level causal labels.",
            "Researcher-view precision and recall are not estimated because that view is an interface policy, not an implemented OCR detector.",
        ],
        "input_sha256": sha256(INPUT_XLSX),
        "outputs": {},
    }
    manifest_path = OUTPUT_DIR / "paired_propagation_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for output in [page_path, summary_path, prediction_path]:
        manifest["outputs"][output.name] = sha256(output)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    lines = [
        "# Paired OCR propagation sensitivity",
        "",
        "This analysis uses the 17 pages with raw OCR, cleaned OCR, and an adjudicated manual transcription.",
        "The same frozen model is run on all three versions. The adjudicated-text prediction is used only as a stability reference, not as human-annotated truth.",
        "",
    ]
    for row in all_summary_rows:
        lines.append(
            f"- {row['model']} / {row['output_level']}: raw stability F1={row['raw_stability_f1']:.4f}; "
            f"cleaned stability F1={row['cleaned_stability_f1']:.4f}; delta={row['cleaned_minus_raw_stability_f1']:.4f} "
            f"(bootstrap 95% CI {row['bootstrap_ci_low']:.4f} to {row['bootstrap_ci_high']:.4f})."
        )
    lines.extend(
        [
            "",
            "These are paired output-stability results. They do not estimate human-reference accuracy, the number of OCR-caused false inferences in the full graph, C-class flagging precision, or researcher-view OCR-detection precision/recall.",
            "",
        ]
    )
    summary_md = OUTPUT_DIR / "paired_propagation_summary.md"
    summary_md.write_text("\n".join(lines), encoding="utf-8")
    print(summary_md.read_text(encoding="utf-8"), flush=True)


if __name__ == "__main__":
    main()
