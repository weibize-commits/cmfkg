from __future__ import annotations

import csv
import gzip
import hashlib
import json
import random
from collections import defaultdict
from copy import copy
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


SCRIPT = Path(__file__).resolve()
RULE_VALIDATION_ROOT = SCRIPT.parents[1]
PAPER_REVISION_ROOT = SCRIPT.parents[2]
WORKSPACE_ROOT = PAPER_REVISION_ROOT.parent

KG_PATH = WORKSPACE_ROOT / "kg_triples_final_v3.jsonl"
CHUNK_PATH = WORKSPACE_ROOT / "kg_canonical.jsonl"
OUTPUT_DIR = RULE_VALIDATION_ROOT / "outputs"
FORM_DIR = RULE_VALIDATION_ROOT / "forms"

SEED = 20260731
SAMPLE_PER_RULE = 100

RULES = {
    "B1": {
        "name_en": "ingredient-meridian-organ-disease candidate",
        "name_zh": "食材—归经—脏腑—病症候选",
        "expression": (
            "FoodIngredient(i) ∧ entersmeridian(i,m) ∧ "
            "correspondsToOrgan(m,o) ∧ DiseaseSymptom(d) ∧ "
            "affectsOrgan(d,o) → candidateTreats(i,d)"
        ),
    },
    "B2": {
        "name_en": "recipe historical-use candidate propagation",
        "name_zh": "食疗方继承食材历史治疗关系候选",
        "expression": (
            "FoodTherapyRecipe(r) ∧ containsIngredient(r,i) ∧ "
            "treats(i,d) → candidateTreats(r,d)"
        ),
    },
    "B3": {
        "name_en": "recipe-ingredient reverse candidate propagation",
        "name_zh": "病症—食疗方—食材反向候选",
        "expression": (
            "DiseaseSymptom(d) ∧ treatedByRecipe(d,r) ∧ "
            "containsIngredient(r,i) → candidateTreats(i,d)"
        ),
    },
}

AUTHORITY_RANK = {
    "single": 1,
    "medium": 2,
    "high": 3,
    "very_high": 4,
}

DECISIONS = ["支持", "不支持", "不确定"]
REASON_CODES = [
    "无明显问题",
    "原文证据不足",
    "OCR或文本错误",
    "实体边界或类型错误",
    "前件关系不成立",
    "规则逻辑过宽",
    "候选结论与原文矛盾",
    "需要更多上下文",
    "其他",
]
YES_NO = ["否", "是"]

NAVY = "17365D"
BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
ORANGE = "FCE4D6"
RED = "F4CCCC"
GREY = "E7E6E6"
WHITE = "FFFFFF"

THIN_GREY = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN_GREY, right=THIN_GREY, top=THIN_GREY, bottom=THIN_GREY)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def stable_id(rule_id: str, head: str, tail: str) -> str:
    raw = f"{rule_id}\t{head}\tcandidateTreats\t{tail}".encode("utf-8")
    return f"{rule_id}-{hashlib.sha256(raw).hexdigest()[:12].upper()}"


def edge_key(edge: dict[str, Any]) -> tuple[str, str]:
    return edge["head"]["text"], edge["tail"]["text"]


def edge_score(edge: dict[str, Any]) -> tuple[float, ...]:
    return (
        float(AUTHORITY_RANK.get(edge.get("authority", ""), 0)),
        float(edge.get("unique_book_count", 0) or 0),
        float(edge.get("support_count", 0) or 0),
        float(edge.get("confidence", 0) or 0),
    )


def best_edges(edges: Iterable[dict[str, Any]]) -> dict[tuple[str, str], dict[str, Any]]:
    selected: dict[tuple[str, str], dict[str, Any]] = {}
    for edge in edges:
        key = edge_key(edge)
        if key not in selected or edge_score(edge) > edge_score(selected[key]):
            selected[key] = edge
    return selected


def load_relations() -> dict[str, list[dict[str, Any]]]:
    wanted = {
        "entersmeridian",
        "correspondsToOrgan",
        "affectsOrgan",
        "containsIngredient",
        "treats",
        "treatedByRecipe",
    }
    relations: dict[str, list[dict[str, Any]]] = defaultdict(list)
    with KG_PATH.open(encoding="utf-8") as handle:
        for line in handle:
            row = json.loads(line)
            if row.get("relation") in wanted:
                relations[row["relation"]].append(row)
    return relations


def valid_signature(
    edge: dict[str, Any],
    head_type: str,
    tail_type: str,
) -> bool:
    return (
        edge.get("head", {}).get("type") == head_type
        and edge.get("tail", {}).get("type") == tail_type
    )


def path_score(edges: tuple[dict[str, Any], ...]) -> tuple[float, ...]:
    ranks = [AUTHORITY_RANK.get(edge.get("authority", ""), 0) for edge in edges]
    supports = [int(edge.get("support_count", 0) or 0) for edge in edges]
    books = [int(edge.get("unique_book_count", 0) or 0) for edge in edges]
    confidences = [float(edge.get("confidence", 0) or 0) for edge in edges]
    return (
        float(min(ranks, default=0)),
        float(sum(ranks)),
        float(sum(books)),
        float(sum(supports)),
        float(sum(confidences)),
    )


def update_candidate(
    candidates: dict[tuple[str, str], dict[str, Any]],
    rule_id: str,
    head: str,
    tail: str,
    path_edges: tuple[dict[str, Any], ...],
) -> None:
    key = (head, tail)
    score = path_score(path_edges)
    if key not in candidates:
        candidates[key] = {
            "candidate_id": stable_id(rule_id, head, tail),
            "rule_id": rule_id,
            "rule_name_en": RULES[rule_id]["name_en"],
            "rule_name_zh": RULES[rule_id]["name_zh"],
            "rule_expression": RULES[rule_id]["expression"],
            "head": head,
            "relation": "candidateTreats",
            "tail": tail,
            "path_count": 1,
            "best_path_score": score,
            "path_edges": path_edges,
        }
        return
    candidates[key]["path_count"] += 1
    if score > candidates[key]["best_path_score"]:
        candidates[key]["best_path_score"] = score
        candidates[key]["path_edges"] = path_edges


def generate_candidates(
    relations: dict[str, list[dict[str, Any]]],
) -> dict[str, list[dict[str, Any]]]:
    enters = best_edges(
        edge
        for edge in relations["entersmeridian"]
        if valid_signature(edge, "食材", "归经")
    )
    corresponds = best_edges(
        edge
        for edge in relations["correspondsToOrgan"]
        if valid_signature(edge, "归经", "脏腑")
    )
    affects = best_edges(
        edge
        for edge in relations["affectsOrgan"]
        if valid_signature(edge, "病症", "脏腑")
    )
    contains = best_edges(
        edge
        for edge in relations["containsIngredient"]
        if valid_signature(edge, "食疗方", "食材")
    )
    treats = best_edges(
        edge
        for edge in relations["treats"]
        if valid_signature(edge, "食材", "病症")
    )
    treated_by = best_edges(
        edge
        for edge in relations["treatedByRecipe"]
        if valid_signature(edge, "病症", "食疗方")
    )

    asserted_treats = {
        edge_key(edge)
        for edge in relations["treats"]
    }

    m_to_correspondence: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for (meridian, _organ), edge in corresponds.items():
        m_to_correspondence[meridian].append(edge)
    organ_to_affects: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for (_disease, organ), edge in affects.items():
        organ_to_affects[organ].append(edge)
    ingredient_to_treats: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for (ingredient, _disease), edge in treats.items():
        ingredient_to_treats[ingredient].append(edge)
    recipe_to_treated_by: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for (_disease, recipe), edge in treated_by.items():
        recipe_to_treated_by[recipe].append(edge)

    by_rule: dict[str, dict[tuple[str, str], dict[str, Any]]] = {
        "B1": {},
        "B2": {},
        "B3": {},
    }

    for (ingredient, meridian), edge_1 in enters.items():
        for edge_2 in m_to_correspondence.get(meridian, []):
            organ = edge_2["tail"]["text"]
            for edge_3 in organ_to_affects.get(organ, []):
                disease = edge_3["head"]["text"]
                if (ingredient, disease) in asserted_treats:
                    continue
                update_candidate(
                    by_rule["B1"],
                    "B1",
                    ingredient,
                    disease,
                    (edge_1, edge_2, edge_3),
                )

    for (recipe, ingredient), edge_1 in contains.items():
        for edge_2 in ingredient_to_treats.get(ingredient, []):
            disease = edge_2["tail"]["text"]
            if (recipe, disease) in asserted_treats:
                continue
            update_candidate(
                by_rule["B2"],
                "B2",
                recipe,
                disease,
                (edge_1, edge_2),
            )

    for (recipe, ingredient), edge_2 in contains.items():
        for edge_1 in recipe_to_treated_by.get(recipe, []):
            disease = edge_1["head"]["text"]
            if (ingredient, disease) in asserted_treats:
                continue
            update_candidate(
                by_rule["B3"],
                "B3",
                ingredient,
                disease,
                (edge_1, edge_2),
            )

    final: dict[str, list[dict[str, Any]]] = {}
    for rule_id, candidates in by_rule.items():
        final[rule_id] = sorted(
            candidates.values(),
            key=lambda row: (row["head"], row["tail"], row["candidate_id"]),
        )
    return final


def edge_compact(edge: dict[str, Any]) -> dict[str, Any]:
    return {
        "head": edge["head"]["text"],
        "head_type": edge["head"].get("type", ""),
        "relation": edge["relation"],
        "tail": edge["tail"]["text"],
        "tail_type": edge["tail"].get("type", ""),
        "authority": edge.get("authority", ""),
        "confidence": edge.get("confidence", ""),
        "support_count": edge.get("support_count", ""),
        "unique_book_count": edge.get("unique_book_count", ""),
        "source_file": edge.get("best_book", ""),
        "source_chunk": edge.get("best_chunk", ""),
        "evidence": edge.get("evidence", ""),
    }


def candidate_serialisable(candidate: dict[str, Any]) -> dict[str, Any]:
    return {
        "candidate_id": candidate["candidate_id"],
        "rule_id": candidate["rule_id"],
        "rule_name_en": candidate["rule_name_en"],
        "rule_name_zh": candidate["rule_name_zh"],
        "rule_expression": candidate["rule_expression"],
        "head": candidate["head"],
        "relation": candidate["relation"],
        "tail": candidate["tail"],
        "path_count": candidate["path_count"],
        "selected_path": [
            edge_compact(edge)
            for edge in candidate["path_edges"]
        ],
        "status": "candidate_not_asserted_fact",
    }


def write_candidate_universe(
    by_rule: dict[str, list[dict[str, Any]]],
) -> tuple[Path, Path]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    csv_path = OUTPUT_DIR / "b_rule_candidate_universe.csv"
    jsonl_path = OUTPUT_DIR / "b_rule_candidate_provenance.jsonl.gz"

    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "candidate_id",
                "rule_id",
                "head",
                "relation",
                "tail",
                "path_count",
                "selected_path_min_authority",
                "selected_path_source_files",
                "status",
            ],
        )
        writer.writeheader()
        for rule_id in RULES:
            for candidate in by_rule[rule_id]:
                edges = candidate["path_edges"]
                min_rank = min(
                    AUTHORITY_RANK.get(edge.get("authority", ""), 0)
                    for edge in edges
                )
                rank_to_name = {value: key for key, value in AUTHORITY_RANK.items()}
                writer.writerow(
                    {
                        "candidate_id": candidate["candidate_id"],
                        "rule_id": rule_id,
                        "head": candidate["head"],
                        "relation": candidate["relation"],
                        "tail": candidate["tail"],
                        "path_count": candidate["path_count"],
                        "selected_path_min_authority": rank_to_name.get(min_rank, ""),
                        "selected_path_source_files": " | ".join(
                            dict.fromkeys(
                                str(edge.get("best_book", ""))
                                for edge in edges
                                if edge.get("best_book")
                            )
                        ),
                        "status": "candidate_not_asserted_fact",
                    }
                )

    with gzip.open(jsonl_path, "wt", encoding="utf-8") as handle:
        for rule_id in RULES:
            for candidate in by_rule[rule_id]:
                handle.write(
                    json.dumps(
                        candidate_serialisable(candidate),
                        ensure_ascii=False,
                    )
                    + "\n"
                )
    return csv_path, jsonl_path


def deterministic_sample(
    by_rule: dict[str, list[dict[str, Any]]],
) -> list[dict[str, Any]]:
    sampled: list[dict[str, Any]] = []
    for offset, rule_id in enumerate(RULES, start=1):
        population = by_rule[rule_id]
        if len(population) < SAMPLE_PER_RULE:
            raise RuntimeError(
                f"{rule_id} has only {len(population)} candidates; "
                f"cannot sample {SAMPLE_PER_RULE}"
            )
        rng = random.Random(SEED + offset)
        indices = sorted(rng.sample(range(len(population)), SAMPLE_PER_RULE))
        sampled.extend(copy(population[index]) for index in indices)
    return sampled


def required_chunk_keys(sample: list[dict[str, Any]]) -> set[tuple[str, int]]:
    keys: set[tuple[str, int]] = set()
    for candidate in sample:
        for edge in candidate["path_edges"]:
            book = edge.get("best_book")
            chunk = edge.get("best_chunk")
            if book and isinstance(chunk, int):
                keys.add((str(book), chunk))
    return keys


def load_contexts(keys: set[tuple[str, int]]) -> dict[tuple[str, int], str]:
    contexts: dict[tuple[str, int], str] = {}
    if not CHUNK_PATH.exists():
        return contexts
    with CHUNK_PATH.open(encoding="utf-8") as handle:
        for line in handle:
            row = json.loads(line)
            key = (str(row.get("book", "")), int(row.get("chunk_idx", -1)))
            if key in keys:
                contexts[key] = str(row.get("text", ""))
                if len(contexts) == len(keys):
                    break
    return contexts


def context_window(text: str, evidence: str, width: int = 900) -> str:
    clean = " ".join(str(text or "").split())
    evidence_clean = " ".join(str(evidence or "").split())
    if not clean:
        return evidence_clean
    if len(clean) <= width:
        return clean
    position = clean.find(evidence_clean) if evidence_clean else -1
    if position < 0:
        return clean[:width] + "…"
    left = max(0, position - width // 2)
    right = min(len(clean), left + width)
    prefix = "…" if left else ""
    suffix = "…" if right < len(clean) else ""
    return prefix + clean[left:right] + suffix


def enrich_sample_context(
    sample: list[dict[str, Any]],
    contexts: dict[tuple[str, int], str],
) -> None:
    for candidate in sample:
        for edge in candidate["path_edges"]:
            key = (
                str(edge.get("best_book", "")),
                int(edge.get("best_chunk", -1)),
            )
            edge["_context"] = context_window(
                contexts.get(key, ""),
                str(edge.get("evidence", "")),
            )


def flatten_sample(candidate: dict[str, Any]) -> dict[str, Any]:
    edges = candidate["path_edges"]
    row: dict[str, Any] = {
        "candidate_id": candidate["candidate_id"],
        "rule_id": candidate["rule_id"],
        "rule_name_zh": candidate["rule_name_zh"],
        "rule_expression": candidate["rule_expression"],
        "head": candidate["head"],
        "relation": candidate["relation"],
        "tail": candidate["tail"],
        "path_count": candidate["path_count"],
        "candidate_statement": (
            f"{candidate['head']} —候选治疗关系→ {candidate['tail']}"
        ),
    }
    for index in range(3):
        prefix = f"edge{index + 1}_"
        if index >= len(edges):
            for field in [
                "triple",
                "source_file",
                "source_chunk",
                "evidence",
                "context",
                "authority",
                "support_count",
            ]:
                row[prefix + field] = ""
            continue
        edge = edges[index]
        row[prefix + "triple"] = (
            f"{edge['head']['text']} | {edge['relation']} | "
            f"{edge['tail']['text']}"
        )
        row[prefix + "source_file"] = edge.get("best_book", "")
        row[prefix + "source_chunk"] = edge.get("best_chunk", "")
        row[prefix + "evidence"] = edge.get("evidence", "")
        row[prefix + "context"] = edge.get("_context", "")
        row[prefix + "authority"] = edge.get("authority", "")
        row[prefix + "support_count"] = edge.get("support_count", "")
    return row


def write_sample_csv(sample_rows: list[dict[str, Any]]) -> Path:
    path = OUTPUT_DIR / "b_rule_human_validation_sample_300.csv"
    fieldnames = list(sample_rows[0].keys())
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(sample_rows)
    return path


def style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = BORDER


def style_body(ws, start_row: int, end_row: int, input_start: int | None = None) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=end_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if input_start and cell.column >= input_start:
                cell.fill = PatternFill("solid", fgColor=GREEN)


def set_widths(ws, widths: dict[int, float]) -> None:
    for column_index, width in widths.items():
        ws.column_dimensions[get_column_letter(column_index)].width = width


def add_instruction_sheet(wb: Workbook, *, adjudicator: bool = False) -> None:
    ws = wb.active
    ws.title = "填写说明"
    title = "CMFKG B1–B3规则人工验证"
    if adjudicator:
        title += "：第三专家仲裁表"
    else:
        title += "：独立盲评表"
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells("A1:F1")

    instructions = [
        ("研究目的", "评价规则产生的候选关系，不评价模型输出；候选关系不是已确认事实。"),
        ("独立性", "专家1和专家2必须独立填写，不得查看或讨论对方的结论。第三专家仅处理分歧项。"),
        ("主要结论", "支持：现有原文证据和完整规则路径足以支持该候选关系。"),
        ("不支持", "任一关键前件不成立、规则推理过宽、实体/OCR错误，或候选结论与原文冲突。"),
        ("不确定", "现有上下文不足以作出可靠判断；不得为了减少缺失而强行归入支持。"),
        ("主要统计量", "每条规则的严格支持率；“不确定”不计为支持，同时单独报告。"),
        ("置信度", "1=很低，2=较低，3=中等，4=较高，5=很高。"),
        ("理由", "请选择一个主要理由代码，并填写能够被第三专家复核的简短文字理由。"),
        ("来源核对", "如片段不足，请将“是否需要查看原书”选为“是”，不要凭常识补足原文证据。"),
        ("禁止事项", "不得改动候选编号、规则路径、来源和证据列；不得填写虚构或模型生成的专家判断。"),
    ]
    start = 3
    for offset, (label, text) in enumerate(instructions):
        row = start + offset
        ws.cell(row, 1, label)
        ws.cell(row, 2, text)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        ws.cell(row, 1).font = Font(bold=True)
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        for cell in ws[row]:
            cell.border = BORDER

    ws.cell(start + len(instructions) + 1, 1, "抽样设计")
    ws.cell(
        start + len(instructions) + 1,
        2,
        (
            f"按规则分层，每条规则从修正后且未被原KG直接断言的候选中随机抽取"
            f"{SAMPLE_PER_RULE}条；总计300条；随机种子{SEED}。"
        ),
    )
    ws.merge_cells(
        start_row=start + len(instructions) + 1,
        start_column=2,
        end_row=start + len(instructions) + 1,
        end_column=6,
    )
    ws.cell(start + len(instructions) + 1, 1).font = Font(bold=True)
    ws.cell(start + len(instructions) + 1, 1).fill = PatternFill("solid", fgColor=YELLOW)
    for cell in ws[start + len(instructions) + 1]:
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    set_widths(ws, {1: 20, 2: 24, 3: 18, 4: 18, 5: 18, 6: 18})
    ws.sheet_view.showGridLines = False


def add_expert_info_sheet(wb: Workbook, role: str) -> None:
    ws = wb.create_sheet("专家信息")
    rows = [
        ("字段", "填写内容"),
        ("专家匿名编号", ""),
        ("角色", role),
        ("最高学位", ""),
        ("专业方向", ""),
        ("中医药/古籍相关工作年限", ""),
        ("当前职称或岗位", ""),
        ("是否参与CMFKG规则设计", ""),
        ("是否存在利益冲突", ""),
        ("培训材料版本", "B-rule-human-validation-v1.0"),
        ("开始日期", ""),
        ("完成日期", ""),
        ("电子确认", "本人确认以上判断为独立完成，未查看其他专家的答案。"),
    ]
    for row in rows:
        ws.append(row)
    style_header(ws)
    style_body(ws, 2, len(rows), input_start=2)
    for row in [3, 10, 13]:
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=GREY)
    set_widths(ws, {1: 30, 2: 80})
    ws.freeze_panes = "A2"


RATING_HEADERS = [
    "盲评序号",
    "候选编号",
    "规则编号",
    "规则中文名称",
    "规则表达式",
    "候选头实体",
    "候选关系",
    "候选尾实体",
    "候选表述",
    "候选路径数量",
    "前件1三元组",
    "前件1来源文件",
    "前件1分块",
    "前件1证据",
    "前件1上下文",
    "前件2三元组",
    "前件2来源文件",
    "前件2分块",
    "前件2证据",
    "前件2上下文",
    "前件3三元组",
    "前件3来源文件",
    "前件3分块",
    "前件3证据",
    "前件3上下文",
    "专家结论*",
    "置信度1-5*",
    "主要理由代码*",
    "文字理由*",
    "是否需要查看原书",
    "填写日期",
    "专家匿名编号",
]


def rating_row(flat: dict[str, Any], order: int) -> list[Any]:
    return [
        order,
        flat["candidate_id"],
        flat["rule_id"],
        flat["rule_name_zh"],
        flat["rule_expression"],
        flat["head"],
        flat["relation"],
        flat["tail"],
        flat["candidate_statement"],
        flat["path_count"],
        flat["edge1_triple"],
        flat["edge1_source_file"],
        flat["edge1_source_chunk"],
        flat["edge1_evidence"],
        flat["edge1_context"],
        flat["edge2_triple"],
        flat["edge2_source_file"],
        flat["edge2_source_chunk"],
        flat["edge2_evidence"],
        flat["edge2_context"],
        flat["edge3_triple"],
        flat["edge3_source_file"],
        flat["edge3_source_chunk"],
        flat["edge3_evidence"],
        flat["edge3_context"],
        "",
        "",
        "",
        "",
        "",
        "",
        "",
    ]


def add_rating_sheet(
    wb: Workbook,
    sample_rows: list[dict[str, Any]],
    expert_label: str,
) -> None:
    ws = wb.create_sheet("盲评")
    ws.append(RATING_HEADERS)
    ordered = list(sample_rows)
    random.Random(f"{SEED}-{expert_label}").shuffle(ordered)
    for index, flat in enumerate(ordered, start=1):
        ws.append(rating_row(flat, index))

    style_header(ws)
    style_body(ws, 2, ws.max_row, input_start=26)
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    decision_dv = DataValidation(
        type="list",
        formula1='"' + ",".join(DECISIONS) + '"',
        allow_blank=True,
    )
    confidence_dv = DataValidation(
        type="whole",
        operator="between",
        formula1="1",
        formula2="5",
        allow_blank=True,
    )
    reason_dv = DataValidation(
        type="list",
        formula1='"' + ",".join(REASON_CODES) + '"',
        allow_blank=True,
    )
    yes_no_dv = DataValidation(
        type="list",
        formula1='"' + ",".join(YES_NO) + '"',
        allow_blank=True,
    )
    ws.add_data_validation(decision_dv)
    ws.add_data_validation(confidence_dv)
    ws.add_data_validation(reason_dv)
    ws.add_data_validation(yes_no_dv)
    decision_dv.add(f"Z2:Z{ws.max_row}")
    confidence_dv.add(f"AA2:AA{ws.max_row}")
    reason_dv.add(f"AB2:AB{ws.max_row}")
    yes_no_dv.add(f"AD2:AD{ws.max_row}")

    missing_fill = PatternFill("solid", fgColor=RED)
    ws.conditional_formatting.add(
        f"Z2:Z{ws.max_row}",
        FormulaRule(formula=["LEN(Z2)=0"], fill=missing_fill),
    )
    ws.conditional_formatting.add(
        f"AA2:AA{ws.max_row}",
        FormulaRule(formula=["LEN(AA2)=0"], fill=missing_fill),
    )
    ws.conditional_formatting.add(
        f"AB2:AB{ws.max_row}",
        FormulaRule(formula=["LEN(AB2)=0"], fill=missing_fill),
    )
    ws.conditional_formatting.add(
        f"AC2:AC{ws.max_row}",
        FormulaRule(formula=["LEN(AC2)=0"], fill=missing_fill),
    )

    set_widths(
        ws,
        {
            1: 10,
            2: 20,
            3: 9,
            4: 24,
            5: 52,
            6: 18,
            7: 18,
            8: 18,
            9: 34,
            10: 12,
            11: 34,
            12: 28,
            13: 10,
            14: 36,
            15: 70,
            16: 34,
            17: 28,
            18: 10,
            19: 36,
            20: 70,
            21: 34,
            22: 28,
            23: 10,
            24: 36,
            25: 70,
            26: 14,
            27: 14,
            28: 24,
            29: 50,
            30: 18,
            31: 14,
            32: 18,
        },
    )
    ws.row_dimensions[1].height = 42
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 120


def add_codebook_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("代码表")
    ws.append(["类别", "代码", "定义"])
    for decision in DECISIONS:
        definitions = {
            "支持": "现有原文证据与规则路径足以支持该候选关系。",
            "不支持": "至少一个关键前件或候选结论不能由现有证据支持。",
            "不确定": "上下文或来源不足，无法作出可靠判断。",
        }
        ws.append(["专家结论", decision, definitions[decision]])
    for index, code in enumerate(REASON_CODES, start=1):
        ws.append(["主要理由", code, "选择最主要的一项，并在文字理由中说明。"])
    for value in range(1, 6):
        ws.append(["置信度", value, f"{value}/5"])
    style_header(ws)
    style_body(ws, 2, ws.max_row)
    set_widths(ws, {1: 18, 2: 26, 3: 80})


def add_sampling_sheet(
    wb: Workbook,
    by_rule: dict[str, list[dict[str, Any]]],
) -> None:
    ws = wb.create_sheet("抽样设计")
    ws.append(
        [
            "规则",
            "修正后候选总体",
            "抽样数",
            "抽样比例",
            "随机种子",
            "抽样单位",
            "主要估计量",
        ]
    )
    for offset, rule_id in enumerate(RULES, start=1):
        population = len(by_rule[rule_id])
        ws.append(
            [
                rule_id,
                population,
                SAMPLE_PER_RULE,
                SAMPLE_PER_RULE / population,
                SEED + offset,
                "唯一候选关系（规则×头实体×尾实体）",
                "严格支持率及Wilson 95%置信区间",
            ]
        )
    style_header(ws)
    style_body(ws, 2, ws.max_row)
    for cell in ws["D"][1:]:
        cell.number_format = "0.000%"
    set_widths(ws, {1: 10, 2: 18, 3: 12, 4: 14, 5: 14, 6: 36, 7: 42})


def add_adjudication_sheet(
    wb: Workbook,
    sample_rows: list[dict[str, Any]],
) -> None:
    headers = [
        "候选编号",
        "规则编号",
        "规则表达式",
        "候选头实体",
        "候选关系",
        "候选尾实体",
        "候选表述",
        "候选路径数量",
        "前件1三元组",
        "前件1来源文件",
        "前件1分块",
        "前件1证据",
        "前件1上下文",
        "前件2三元组",
        "前件2来源文件",
        "前件2分块",
        "前件2证据",
        "前件2上下文",
        "前件3三元组",
        "前件3来源文件",
        "前件3分块",
        "前件3证据",
        "前件3上下文",
        "专家1结论",
        "专家1置信度",
        "专家1理由代码",
        "专家1理由",
        "专家2结论",
        "专家2置信度",
        "专家2理由代码",
        "专家2理由",
        "最终结论*",
        "最终置信度1-5*",
        "仲裁理由代码*",
        "仲裁文字理由*",
        "仲裁日期",
        "仲裁专家匿名编号",
    ]
    ws = wb.create_sheet("仲裁")
    ws.append(headers)
    for flat in sorted(sample_rows, key=lambda row: row["candidate_id"]):
        ws.append(
            [
                flat["candidate_id"],
                flat["rule_id"],
                flat["rule_expression"],
                flat["head"],
                flat["relation"],
                flat["tail"],
                flat["candidate_statement"],
                flat["path_count"],
                flat["edge1_triple"],
                flat["edge1_source_file"],
                flat["edge1_source_chunk"],
                flat["edge1_evidence"],
                flat["edge1_context"],
                flat["edge2_triple"],
                flat["edge2_source_file"],
                flat["edge2_source_chunk"],
                flat["edge2_evidence"],
                flat["edge2_context"],
                flat["edge3_triple"],
                flat["edge3_source_file"],
                flat["edge3_source_chunk"],
                flat["edge3_evidence"],
                flat["edge3_context"],
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
            ]
        )
    style_header(ws)
    style_body(ws, 2, ws.max_row, input_start=32)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    decision_dv = DataValidation(
        type="list",
        formula1='"' + ",".join(DECISIONS) + '"',
        allow_blank=True,
    )
    confidence_dv = DataValidation(
        type="whole",
        operator="between",
        formula1="1",
        formula2="5",
        allow_blank=True,
    )
    reason_dv = DataValidation(
        type="list",
        formula1='"' + ",".join(REASON_CODES) + '"',
        allow_blank=True,
    )
    ws.add_data_validation(decision_dv)
    ws.add_data_validation(confidence_dv)
    ws.add_data_validation(reason_dv)
    decision_dv.add(f"AF2:AF{ws.max_row}")
    confidence_dv.add(f"AG2:AG{ws.max_row}")
    reason_dv.add(f"AH2:AH{ws.max_row}")
    set_widths(
        ws,
        {
            1: 20,
            2: 10,
            3: 55,
            4: 18,
            5: 18,
            6: 18,
            7: 35,
            8: 12,
            9: 34,
            10: 28,
            11: 10,
            12: 36,
            13: 70,
            14: 34,
            15: 28,
            16: 10,
            17: 36,
            18: 70,
            19: 34,
            20: 28,
            21: 10,
            22: 36,
            23: 70,
            24: 14,
            25: 14,
            26: 24,
            27: 45,
            28: 14,
            29: 14,
            30: 24,
            31: 45,
            32: 14,
            33: 16,
            34: 24,
            35: 50,
            36: 14,
            37: 20,
        },
    )
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 120


def create_expert_workbook(
    path: Path,
    sample_rows: list[dict[str, Any]],
    expert_label: str,
) -> None:
    wb = Workbook()
    add_instruction_sheet(wb)
    add_expert_info_sheet(wb, f"{expert_label}独立盲评")
    add_rating_sheet(wb, sample_rows, expert_label)
    add_codebook_sheet(wb)
    wb.save(path)


def create_adjudication_workbook(
    path: Path,
    sample_rows: list[dict[str, Any]],
) -> None:
    wb = Workbook()
    add_instruction_sheet(wb, adjudicator=True)
    add_expert_info_sheet(wb, "第三专家分歧仲裁")
    add_adjudication_sheet(wb, sample_rows)
    add_codebook_sheet(wb)
    wb.save(path)


def copy_sheet(source, target) -> None:
    for row in source.iter_rows():
        for cell in row:
            new_cell = target[cell.coordinate]
            new_cell.value = cell.value
            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.fill = copy(cell.fill)
                new_cell.border = copy(cell.border)
                new_cell.alignment = copy(cell.alignment)
                new_cell.number_format = cell.number_format
                new_cell.protection = copy(cell.protection)
    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key].width = dimension.width
        target.column_dimensions[key].hidden = dimension.hidden
    for key, dimension in source.row_dimensions.items():
        target.row_dimensions[key].height = dimension.height
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    if source.auto_filter.ref:
        target.auto_filter.ref = source.auto_filter.ref
    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))


def create_master_workbook(
    path: Path,
    by_rule: dict[str, list[dict[str, Any]]],
    sample_rows: list[dict[str, Any]],
    expert1_path: Path,
    expert2_path: Path,
    adjudication_path: Path,
) -> None:
    wb = Workbook()
    add_instruction_sheet(wb)
    ws = wb["填写说明"]
    next_row = ws.max_row + 3
    ws.cell(next_row, 1, "文件分发")
    ws.cell(
        next_row,
        2,
        (
            "为保持盲评独立性，请分别发送01和02文件；两位专家完成后由协调员合并，"
            "只把分歧项发送给第三专家。不要让专家共用本Master文件。"
        ),
    )
    ws.merge_cells(start_row=next_row, start_column=2, end_row=next_row, end_column=6)
    ws.cell(next_row, 1).font = Font(bold=True)
    ws.cell(next_row, 1).fill = PatternFill("solid", fgColor=ORANGE)
    for cell in ws[next_row]:
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    add_sampling_sheet(wb, by_rule)
    add_codebook_sheet(wb)

    for title, form_path, source_sheet in [
        ("专家1盲评", expert1_path, "盲评"),
        ("专家2盲评", expert2_path, "盲评"),
        ("第三专家仲裁", adjudication_path, "仲裁"),
    ]:
        source_wb = __import__("openpyxl").load_workbook(form_path)
        target = wb.create_sheet(title)
        copy_sheet(source_wb[source_sheet], target)
    wb.save(path)


def write_manifest(
    by_rule: dict[str, list[dict[str, Any]]],
    sample_rows: list[dict[str, Any]],
    generated_files: list[Path],
) -> Path:
    manifest_path = OUTPUT_DIR / "b_rule_validation_manifest.json"
    manifest = {
        "package_version": "1.0.0",
        "generated_date": "2026-07-31",
        "input_kg": {
            "path": str(KG_PATH),
            "sha256": sha256_file(KG_PATH),
        },
        "input_chunks": {
            "path": str(CHUNK_PATH),
            "sha256": sha256_file(CHUNK_PATH),
        },
        "candidate_definition": (
            "Unique rule-specific candidate relation after corrected antecedent "
            "matching; any candidate already directly asserted as treats in the "
            "input KG was excluded."
        ),
        "population_counts": {
            rule_id: len(by_rule[rule_id])
            for rule_id in RULES
        },
        "sample_design": {
            "strata": list(RULES),
            "sample_per_rule": SAMPLE_PER_RULE,
            "total_sample": len(sample_rows),
            "base_seed": SEED,
            "rule_seeds": {
                rule_id: SEED + offset
                for offset, rule_id in enumerate(RULES, start=1)
            },
            "selection": (
                "Python random.Random.sample without replacement within each "
                "rule stratum after deterministic lexical sorting."
            ),
            "primary_endpoint": (
                "Final-adjudicated Supported proportion for each rule, with "
                "Wilson 95% confidence interval. Uncertain is not counted as "
                "Supported and is reported separately."
            ),
        },
        "blinding": (
            "Expert 1 and Expert 2 receive separately shuffled workbooks and "
            "must not access each other's decisions. Expert 3 receives only "
            "disagreements after coordinator merge."
        ),
        "generated_files": {},
    }
    for path in generated_files:
        manifest["generated_files"][str(path.relative_to(RULE_VALIDATION_ROOT))] = {
            "sha256": sha256_file(path),
            "bytes": path.stat().st_size,
        }
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return manifest_path


def write_readme(
    by_rule: dict[str, list[dict[str, Any]]],
) -> Path:
    path = RULE_VALIDATION_ROOT / "README_zh.md"
    text = f"""# CMFKG B1–B3规则人工验证包

## 当前状态

修正后的三条规则已在冻结的 `kg_triples_final_v3.jsonl` 上重新执行，并保留逐规则候选和代表性完整证据路径。候选关系只是待验证的 `candidateTreats`，不是已确认事实，也没有写回原知识图谱。

候选总体：

- B1：{len(by_rule["B1"]):,} 条
- B2：{len(by_rule["B2"]):,} 条
- B3：{len(by_rule["B3"]):,} 条

按规则分层抽样，每条规则随机抽取 {SAMPLE_PER_RULE} 条，共300条；基础随机种子为 {SEED}。

## 文件分发

- `forms/01_B_Rules_Expert1_Blind.xlsx`：只发给专家1。
- `forms/02_B_Rules_Expert2_Blind.xlsx`：只发给专家2。
- `forms/03_B_Rules_Expert3_Adjudication.xlsx`：专家1和专家2全部完成并合并后，只保留分歧项再发给专家3。
- `forms/00_B1-B3_Human_Validation_Master.xlsx`：协调员使用，不发给评审专家。

两位主评专家不得共用Master工作簿，也不得互相查看答案。空白评分格是等待真实专家填写，不是缺失实验数据或模拟数据。

## 判定标准

主要终点为每条规则最终仲裁后的“支持”比例及Wilson 95%置信区间。“不确定”不计为支持，并单独报告。专家必须核对证据片段、来源文件、分块编号和完整前件路径；如上下文不足，应选择“不确定”并标记需要查看原书。

## 可复现性

`outputs/b_rule_candidate_universe.csv` 保存全部逐规则候选摘要；`outputs/b_rule_candidate_provenance.jsonl.gz` 保存每条候选的代表性证据路径；`outputs/b_rule_validation_manifest.json` 保存输入和输出哈希、抽样规则及统计终点。旧的合并 `rules_B1_B2_B3` 输出保持不变，仅用于历史审计。
"""
    path.write_text(text, encoding="utf-8")
    return path


def verify_workbook(path: Path, expected_sheet: str) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=False, data_only=False)
    if expected_sheet not in wb.sheetnames:
        raise RuntimeError(f"{path.name}: missing sheet {expected_sheet}")
    ws = wb[expected_sheet]
    if ws.max_row != 301:
        raise RuntimeError(
            f"{path.name}: expected 301 rows in {expected_sheet}, found {ws.max_row}"
        )
    candidate_ids = [ws.cell(row, 2 if expected_sheet == "盲评" else 1).value for row in range(2, 302)]
    if len(set(candidate_ids)) != 300:
        raise RuntimeError(f"{path.name}: candidate IDs are not unique")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    FORM_DIR.mkdir(parents=True, exist_ok=True)

    relations = load_relations()
    by_rule = generate_candidates(relations)
    universe_csv, provenance_jsonl = write_candidate_universe(by_rule)

    sample = deterministic_sample(by_rule)
    contexts = load_contexts(required_chunk_keys(sample))
    enrich_sample_context(sample, contexts)
    sample_rows = [flatten_sample(candidate) for candidate in sample]
    sample_csv = write_sample_csv(sample_rows)

    expert1_path = FORM_DIR / "01_B_Rules_Expert1_Blind.xlsx"
    expert2_path = FORM_DIR / "02_B_Rules_Expert2_Blind.xlsx"
    adjudication_path = FORM_DIR / "03_B_Rules_Expert3_Adjudication.xlsx"
    master_path = FORM_DIR / "00_B1-B3_Human_Validation_Master.xlsx"

    create_expert_workbook(expert1_path, sample_rows, "Expert1")
    create_expert_workbook(expert2_path, sample_rows, "Expert2")
    create_adjudication_workbook(adjudication_path, sample_rows)
    create_master_workbook(
        master_path,
        by_rule,
        sample_rows,
        expert1_path,
        expert2_path,
        adjudication_path,
    )
    readme_path = write_readme(by_rule)
    manifest_path = write_manifest(
        by_rule,
        sample_rows,
        [
            SCRIPT,
            RULE_VALIDATION_ROOT / "analysis" / "analyze_b_rule_validation.py",
            universe_csv,
            provenance_jsonl,
            sample_csv,
            expert1_path,
            expert2_path,
            adjudication_path,
            master_path,
            readme_path,
        ],
    )

    verify_workbook(expert1_path, "盲评")
    verify_workbook(expert2_path, "盲评")
    verify_workbook(adjudication_path, "仲裁")

    print(
        json.dumps(
            {
                "population_counts": {
                    rule_id: len(by_rule[rule_id])
                    for rule_id in RULES
                },
                "sample_counts": {
                    rule_id: sum(
                        row["rule_id"] == rule_id
                        for row in sample_rows
                    )
                    for rule_id in RULES
                },
                "forms": [
                    str(master_path),
                    str(expert1_path),
                    str(expert2_path),
                    str(adjudication_path),
                ],
                "manifest": str(manifest_path),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
