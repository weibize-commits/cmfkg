from __future__ import annotations

import argparse
import csv
import json
import math
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SCRIPT = Path(__file__).resolve()
ROOT = SCRIPT.parents[1]
FORM_DIR = ROOT / "forms"
OUTPUT_DIR = ROOT / "outputs"

DEFAULT_EXPERT1 = FORM_DIR / "01_B_Rules_Expert1_Blind.xlsx"
DEFAULT_EXPERT2 = FORM_DIR / "02_B_Rules_Expert2_Blind.xlsx"
DEFAULT_ADJUDICATION_TEMPLATE = FORM_DIR / "03_B_Rules_Expert3_Adjudication.xlsx"
DEFAULT_DISAGREEMENT_FORM = FORM_DIR / "04_B_Rules_Expert3_Disagreements.xlsx"

ALLOWED_DECISIONS = {"支持", "不支持", "不确定"}
RATING_SHEET = "盲评"
ADJUDICATION_SHEET = "仲裁"


def header_map(ws) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[1]
        if cell.value is not None
    }


def normalise_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_primary_ratings(path: Path, expert_label: str) -> dict[str, dict[str, Any]]:
    wb = load_workbook(path, read_only=False, data_only=True)
    if RATING_SHEET not in wb.sheetnames:
        raise RuntimeError(f"{path.name}: missing sheet {RATING_SHEET}")
    ws = wb[RATING_SHEET]
    columns = header_map(ws)
    required = [
        "候选编号",
        "规则编号",
        "专家结论*",
        "置信度1-5*",
        "主要理由代码*",
        "文字理由*",
    ]
    missing_headers = [name for name in required if name not in columns]
    if missing_headers:
        raise RuntimeError(f"{path.name}: missing columns {missing_headers}")

    rows: dict[str, dict[str, Any]] = {}
    incomplete: list[str] = []
    for row_index in range(2, ws.max_row + 1):
        candidate_id = normalise_text(
            ws.cell(row_index, columns["候选编号"]).value
        )
        if not candidate_id:
            continue
        if candidate_id in rows:
            raise RuntimeError(f"{path.name}: duplicate candidate {candidate_id}")
        decision = normalise_text(
            ws.cell(row_index, columns["专家结论*"]).value
        )
        confidence = ws.cell(row_index, columns["置信度1-5*"]).value
        reason_code = normalise_text(
            ws.cell(row_index, columns["主要理由代码*"]).value
        )
        rationale = normalise_text(
            ws.cell(row_index, columns["文字理由*"]).value
        )
        if (
            decision not in ALLOWED_DECISIONS
            or confidence not in {1, 2, 3, 4, 5}
            or not reason_code
            or not rationale
        ):
            incomplete.append(candidate_id)
        rows[candidate_id] = {
            "candidate_id": candidate_id,
            "rule_id": normalise_text(
                ws.cell(row_index, columns["规则编号"]).value
            ),
            "decision": decision,
            "confidence": confidence,
            "reason_code": reason_code,
            "rationale": rationale,
            "expert": expert_label,
        }

    if len(rows) != 300:
        raise RuntimeError(f"{path.name}: expected 300 candidates, found {len(rows)}")
    if incomplete:
        preview = ", ".join(incomplete[:10])
        raise RuntimeError(
            f"{path.name}: {len(incomplete)} incomplete ratings; first: {preview}"
        )
    return rows


def validate_same_candidates(
    expert1: dict[str, dict[str, Any]],
    expert2: dict[str, dict[str, Any]],
) -> None:
    ids1 = set(expert1)
    ids2 = set(expert2)
    if ids1 != ids2:
        raise RuntimeError(
            "Expert files contain different candidates: "
            f"only Expert1={len(ids1 - ids2)}, only Expert2={len(ids2 - ids1)}"
        )
    for candidate_id in ids1:
        if expert1[candidate_id]["rule_id"] != expert2[candidate_id]["rule_id"]:
            raise RuntimeError(f"Rule mismatch for {candidate_id}")


def prepare_disagreement_workbook(
    template_path: Path,
    output_path: Path,
    expert1: dict[str, dict[str, Any]],
    expert2: dict[str, dict[str, Any]],
) -> list[str]:
    wb = load_workbook(template_path, read_only=False, data_only=False)
    if ADJUDICATION_SHEET not in wb.sheetnames:
        raise RuntimeError(
            f"{template_path.name}: missing sheet {ADJUDICATION_SHEET}"
        )
    ws = wb[ADJUDICATION_SHEET]
    columns = header_map(ws)
    required = [
        "候选编号",
        "专家1结论",
        "专家1置信度",
        "专家1理由代码",
        "专家1理由",
        "专家2结论",
        "专家2置信度",
        "专家2理由代码",
        "专家2理由",
    ]
    missing = [name for name in required if name not in columns]
    if missing:
        raise RuntimeError(f"{template_path.name}: missing columns {missing}")

    source_rows: dict[str, list[Any]] = {}
    for row_index in range(2, ws.max_row + 1):
        candidate_id = normalise_text(
            ws.cell(row_index, columns["候选编号"]).value
        )
        if candidate_id:
            source_rows[candidate_id] = [
                ws.cell(row_index, column).value
                for column in range(1, ws.max_column + 1)
            ]
    if len(source_rows) != 300:
        raise RuntimeError(
            f"{template_path.name}: expected 300 template candidates, "
            f"found {len(source_rows)}"
        )

    disagreements = sorted(
        candidate_id
        for candidate_id in expert1
        if expert1[candidate_id]["decision"] != expert2[candidate_id]["decision"]
    )
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for candidate_id in disagreements:
        values = source_rows[candidate_id]
        ws.append(values)
        row_index = ws.max_row
        e1 = expert1[candidate_id]
        e2 = expert2[candidate_id]
        ws.cell(row_index, columns["专家1结论"], e1["decision"])
        ws.cell(row_index, columns["专家1置信度"], e1["confidence"])
        ws.cell(row_index, columns["专家1理由代码"], e1["reason_code"])
        ws.cell(row_index, columns["专家1理由"], e1["rationale"])
        ws.cell(row_index, columns["专家2结论"], e2["decision"])
        ws.cell(row_index, columns["专家2置信度"], e2["confidence"])
        ws.cell(row_index, columns["专家2理由代码"], e2["reason_code"])
        ws.cell(row_index, columns["专家2理由"], e2["rationale"])
    ws.auto_filter.ref = ws.dimensions
    wb.save(output_path)
    return disagreements


def cohen_kappa(pairs: list[tuple[str, str]]) -> float | None:
    if not pairs:
        return None
    total = len(pairs)
    observed = sum(a == b for a, b in pairs) / total
    counts_a = Counter(a for a, _ in pairs)
    counts_b = Counter(b for _, b in pairs)
    expected = sum(
        (counts_a[label] / total) * (counts_b[label] / total)
        for label in ALLOWED_DECISIONS
    )
    if math.isclose(1.0 - expected, 0.0):
        return None
    return (observed - expected) / (1.0 - expected)


def initial_agreement(
    expert1: dict[str, dict[str, Any]],
    expert2: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    result: dict[str, Any] = {}
    groups: dict[str, list[str]] = defaultdict(list)
    groups["all"] = sorted(expert1)
    for candidate_id, rating in expert1.items():
        groups[rating["rule_id"]].append(candidate_id)
    for group, candidate_ids in groups.items():
        pairs = [
            (
                expert1[candidate_id]["decision"],
                expert2[candidate_id]["decision"],
            )
            for candidate_id in candidate_ids
        ]
        exact = sum(a == b for a, b in pairs)
        result[group] = {
            "n": len(pairs),
            "exact_agreement_n": exact,
            "exact_agreement_rate": exact / len(pairs),
            "cohen_kappa_unweighted": cohen_kappa(pairs),
            "expert1_distribution": dict(Counter(a for a, _ in pairs)),
            "expert2_distribution": dict(Counter(b for _, b in pairs)),
        }
    return result


def read_adjudication(path: Path) -> dict[str, dict[str, Any]]:
    wb = load_workbook(path, read_only=False, data_only=True)
    if ADJUDICATION_SHEET not in wb.sheetnames:
        raise RuntimeError(f"{path.name}: missing sheet {ADJUDICATION_SHEET}")
    ws = wb[ADJUDICATION_SHEET]
    columns = header_map(ws)
    required = [
        "候选编号",
        "最终结论*",
        "最终置信度1-5*",
        "仲裁理由代码*",
        "仲裁文字理由*",
    ]
    missing = [name for name in required if name not in columns]
    if missing:
        raise RuntimeError(f"{path.name}: missing columns {missing}")
    ratings: dict[str, dict[str, Any]] = {}
    incomplete: list[str] = []
    for row_index in range(2, ws.max_row + 1):
        candidate_id = normalise_text(
            ws.cell(row_index, columns["候选编号"]).value
        )
        if not candidate_id:
            continue
        decision = normalise_text(
            ws.cell(row_index, columns["最终结论*"]).value
        )
        confidence = ws.cell(row_index, columns["最终置信度1-5*"]).value
        reason_code = normalise_text(
            ws.cell(row_index, columns["仲裁理由代码*"]).value
        )
        rationale = normalise_text(
            ws.cell(row_index, columns["仲裁文字理由*"]).value
        )
        if (
            decision not in ALLOWED_DECISIONS
            or confidence not in {1, 2, 3, 4, 5}
            or not reason_code
            or not rationale
        ):
            incomplete.append(candidate_id)
        ratings[candidate_id] = {
            "decision": decision,
            "confidence": confidence,
            "reason_code": reason_code,
            "rationale": rationale,
        }
    if incomplete:
        preview = ", ".join(incomplete[:10])
        raise RuntimeError(
            f"{path.name}: {len(incomplete)} incomplete adjudications; first: "
            f"{preview}"
        )
    return ratings


def wilson_interval(
    successes: int,
    total: int,
    z: float = 1.959963984540054,
) -> tuple[float, float]:
    if total <= 0:
        return float("nan"), float("nan")
    p = successes / total
    denominator = 1 + z * z / total
    centre = (p + z * z / (2 * total)) / denominator
    half = (
        z
        * math.sqrt(p * (1 - p) / total + z * z / (4 * total * total))
        / denominator
    )
    return centre - half, centre + half


def finalise_items(
    expert1: dict[str, dict[str, Any]],
    expert2: dict[str, dict[str, Any]],
    adjudication: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    final_rows: list[dict[str, Any]] = []
    disagreements = {
        candidate_id
        for candidate_id in expert1
        if expert1[candidate_id]["decision"] != expert2[candidate_id]["decision"]
    }
    missing_adjudications = sorted(disagreements - set(adjudication))
    extra_adjudications = sorted(set(adjudication) - disagreements)
    if missing_adjudications:
        raise RuntimeError(
            f"Missing {len(missing_adjudications)} adjudications; first: "
            + ", ".join(missing_adjudications[:10])
        )
    if extra_adjudications:
        raise RuntimeError(
            f"Adjudication file contains {len(extra_adjudications)} agreements; "
            "only disagreements should be adjudicated."
        )

    for candidate_id in sorted(expert1):
        e1 = expert1[candidate_id]
        e2 = expert2[candidate_id]
        if candidate_id in disagreements:
            final = adjudication[candidate_id]
            decision_source = "Expert3_adjudication"
        else:
            final = {
                "decision": e1["decision"],
                "confidence": "",
                "reason_code": "",
                "rationale": "",
            }
            decision_source = "Expert1_Expert2_agreement"
        final_rows.append(
            {
                "candidate_id": candidate_id,
                "rule_id": e1["rule_id"],
                "expert1_decision": e1["decision"],
                "expert1_confidence": e1["confidence"],
                "expert1_reason_code": e1["reason_code"],
                "expert1_rationale": e1["rationale"],
                "expert2_decision": e2["decision"],
                "expert2_confidence": e2["confidence"],
                "expert2_reason_code": e2["reason_code"],
                "expert2_rationale": e2["rationale"],
                "final_decision": final["decision"],
                "decision_source": decision_source,
                "adjudicator_confidence": final["confidence"],
                "adjudicator_reason_code": final["reason_code"],
                "adjudicator_rationale": final["rationale"],
            }
        )
    return final_rows


def write_final_outputs(
    final_rows: list[dict[str, Any]],
    agreement: dict[str, Any],
) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    item_path = OUTPUT_DIR / "b_rule_item_level_final.csv"
    with item_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(final_rows[0].keys()))
        writer.writeheader()
        writer.writerows(final_rows)

    metric_rows: list[dict[str, Any]] = []
    for rule_id in ["B1", "B2", "B3"]:
        rows = [row for row in final_rows if row["rule_id"] == rule_id]
        decisions = Counter(row["final_decision"] for row in rows)
        supported = decisions["支持"]
        lower, upper = wilson_interval(supported, len(rows))
        metric_rows.append(
            {
                "rule_id": rule_id,
                "n": len(rows),
                "supported_n": supported,
                "not_supported_n": decisions["不支持"],
                "uncertain_n": decisions["不确定"],
                "supported_proportion": supported / len(rows),
                "wilson_95_ci_lower": lower,
                "wilson_95_ci_upper": upper,
                "initial_exact_agreement": agreement[rule_id][
                    "exact_agreement_rate"
                ],
                "initial_cohen_kappa": agreement[rule_id][
                    "cohen_kappa_unweighted"
                ],
            }
        )
    metrics_path = OUTPUT_DIR / "b_rule_human_validation_results.csv"
    with metrics_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(metric_rows[0].keys()))
        writer.writeheader()
        writer.writerows(metric_rows)

    agreement_path = OUTPUT_DIR / "b_rule_interrater_agreement.json"
    agreement_path.write_text(
        json.dumps(agreement, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Validate two completed B-rule expert forms, create a "
            "disagreement-only adjudication workbook, and optionally analyse "
            "a completed adjudication file."
        )
    )
    parser.add_argument("--expert1", type=Path, default=DEFAULT_EXPERT1)
    parser.add_argument("--expert2", type=Path, default=DEFAULT_EXPERT2)
    parser.add_argument(
        "--adjudication-template",
        type=Path,
        default=DEFAULT_ADJUDICATION_TEMPLATE,
    )
    parser.add_argument(
        "--disagreement-output",
        type=Path,
        default=DEFAULT_DISAGREEMENT_FORM,
    )
    parser.add_argument(
        "--completed-adjudication",
        type=Path,
        default=None,
        help="Completed disagreement-only Expert 3 workbook.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    expert1 = read_primary_ratings(args.expert1, "Expert1")
    expert2 = read_primary_ratings(args.expert2, "Expert2")
    validate_same_candidates(expert1, expert2)
    agreement = initial_agreement(expert1, expert2)
    disagreements = prepare_disagreement_workbook(
        args.adjudication_template,
        args.disagreement_output,
        expert1,
        expert2,
    )
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "b_rule_pre_adjudication_agreement.json").write_text(
        json.dumps(agreement, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    result = {
        "expert1_rows": len(expert1),
        "expert2_rows": len(expert2),
        "disagreements": len(disagreements),
        "disagreement_form": str(args.disagreement_output),
        "final_analysis_completed": False,
    }
    if args.completed_adjudication is not None:
        adjudication = read_adjudication(args.completed_adjudication)
        final_rows = finalise_items(expert1, expert2, adjudication)
        write_final_outputs(final_rows, agreement)
        result["final_analysis_completed"] = True
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
